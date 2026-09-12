#!/usr/bin/env python3
"""Independent validation and Excel time-mapping audit for dynamic Q2."""
from __future__ import annotations

from pathlib import Path
from shutil import copy2

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "question2"
N, DT, ETA_C, ETA_D = 144, 1 / 6, .90, .90
START, END = pd.Timestamp("2025-02-01"), pd.Timestamp("2025-12-31")


def slots_ok(df: pd.DataFrame) -> bool:
    return bool((df.groupby("date").size() == N).all() and (
        df.groupby("date")["time_index"].apply(
            lambda x: sorted(x.astype(int)) == list(range(N))
        ).all())
    )


def time_label(minutes: int) -> str:
    day, minute = divmod(minutes, 24 * 60)
    hour, minute = divmod(minute, 60)
    return f"{hour}:{minute:02d}" + (f"+{day}" if day else "")


def compressed_emergency(day: pd.DataFrame) -> list[tuple[str, float]]:
    active = day.loc[day.emergency_purchase_kwh > 1e-8].sort_values("time_index")
    result: list[tuple[str, float]] = []
    for _, group in active.groupby((active.time_index.diff().fillna(1) != 1).cumsum()):
        first, last = int(group.time_index.iloc[0]), int(group.time_index.iloc[-1])
        result.append((f"{time_label(first * 10)}-{time_label((last + 1) * 10)}",
                       float(group.emergency_purchase_kwh.sum())))
    return result


def write_checked_workbook(plan: pd.DataFrame, emergency: pd.DataFrame) -> tuple[Path, bool]:
    """Keep raw workbook untouched; create a corrected cyclic-header copy."""
    raw, checked = OUT / "result2_final.xlsx", OUT / "result2_final_checked.xlsx"
    copy2(raw, checked)
    wb = load_workbook(checked)
    ws = wb["计划购电量"]
    headers = [ws.cell(1, c).value for c in range(2, 146)]
    # The template has minor formatting inconsistencies (for example `7:0`),
    # so establish the semantic cyclic order from its unambiguous endpoints.
    cyclic = len(headers) == N and headers[0] == "0:10-0:20" and headers[-1] == "0:00-0:10+1"
    if not cyclic:
        raise ValueError("Official template headers are not the expected 1..143,0 cycle.")
    dates = pd.date_range(START, END, freq="D")
    for r, date in enumerate(dates, start=2):
        values = plan.loc[plan.date == date].sort_values("time_index").grid_purchase_kwh.to_numpy(float)
        for col_offset, t in enumerate(list(range(1, N)) + [0]):
            ws.cell(r, 2 + col_offset).value = float(values[t])
        ws.cell(r, 146).value = float(values.sum())
        ws.cell(r, 147).value = float(np.dot(values, plan.loc[plan.date == date].sort_values("time_index").price_yuan_per_kwh))
    ws = wb["紧急购电量"]
    blocks = {pd.Timestamp("2025-02-01"): [2, 3, 4], pd.Timestamp("2025-02-02"): [5, 6, 7], pd.Timestamp("2025-12-31"): [9, 10, 11]}
    for date, rows in blocks.items():
        values = compressed_emergency(emergency.loc[emergency.date == date])
        for row in rows:
            ws.cell(row, 2).value = None
            ws.cell(row, 3).value = None
        if not values:
            ws.cell(rows[0], 2).value, ws.cell(rows[0], 3).value = "无", 0.0
        else:
            for row, (period, energy) in zip(rows, values[:len(rows)]):
                ws.cell(row, 2).value, ws.cell(row, 3).value = period, energy
            if len(values) > len(rows):
                ws.cell(rows[-1], 2).value = "；".join(x[0] for x in values[len(rows)-1:])
                ws.cell(rows[-1], 3).value = sum(x[1] for x in values[len(rows)-1:])
    wb.save(checked)
    return checked, cyclic


def main() -> None:
    price = pd.read_csv(ROOT / "data" / "processed" / "attachment1_standard_day.csv")
    actual = pd.read_csv(ROOT / "data" / "processed" / "attachment2_actual_long.csv", parse_dates=["date"])
    forecast = pd.read_csv(OUT / "dynamic_forecasts.csv", parse_dates=["date", "history_end_date"])
    plan = pd.read_csv(OUT / "question2_plan.csv", parse_dates=["date"])
    dispatch = pd.read_csv(OUT / "question2_dispatch.csv", parse_dates=["date"])
    emergency = pd.read_csv(OUT / "question2_emergency.csv", parse_dates=["date"])
    daily = pd.read_csv(OUT / "question2_daily_summary.csv", parse_dates=["date"])
    comparison = pd.read_csv(OUT / "strategy_comparison.csv")
    checks: list[tuple[str, bool, float | None]] = []
    add = lambda name, ok, value=None: checks.append((name, bool(ok), value))
    dates = pd.date_range(START, END, freq="D")
    add("forecast_364_days", len(forecast) == 364 * N and forecast.date.min() == pd.Timestamp("2025-01-02") and forecast.date.max() == END)
    add("forecast_no_future_information", (forecast.history_end_date < forecast.date).all())
    add("formal_result_shapes", len(plan) == len(dispatch) == len(emergency) == 334 * N and len(daily) == 334 and slots_ok(plan) and slots_ok(dispatch) and slots_ok(emergency))
    merged = plan.merge(dispatch, on=["date", "time_index"], validate="one_to_one").merge(emergency, on=["date", "time_index"], validate="one_to_one")
    actual = actual.loc[actual.date.between(START, END), ["date", "time_index", "load_kw", "pv_actual_kw"]]
    merged = merged.merge(actual, on=["date", "time_index"], validate="one_to_one")
    balance = (merged.grid_purchase_kwh + merged.pv_actual_kw * DT + merged.discharge_kwh + merged.emergency_purchase_kwh - merged.load_kw * DT - merged.charge_kwh - merged.spill_kwh)
    max_balance = float(np.abs(balance).max())
    transition = merged.soc_end_kwh - merged.soc_start_kwh - ETA_C * merged.charge_kwh + merged.discharge_kwh / ETA_D
    max_transition = float(np.abs(transition).max())
    continuity = merged.soc_end_kwh.iloc[:-1].to_numpy() - merged.soc_start_kwh.iloc[1:].to_numpy()
    max_continuity = float(np.abs(continuity).max())
    add("energy_balance", max_balance < 1e-6, max_balance)
    add("soc_transition", max_transition < 1e-6, max_transition)
    add("soc_continuity", max_continuity < 1e-6, max_continuity)
    add("soc_bounds", ((merged[["soc_start_kwh", "soc_end_kwh"]] >= 1200 - 1e-6) & (merged[["soc_start_kwh", "soc_end_kwh"]] <= 10800 + 1e-6)).all().all())
    total_cd = merged.charge_kwh + merged.discharge_kwh
    add("charge_discharge_limits", (merged.charge_kwh >= -1e-8).all() and (merged.discharge_kwh >= -1e-8).all() and (total_cd <= 5000 * DT + 1e-6).all())
    simultaneous = int(((merged.charge_kwh > 1e-8) & (merged.discharge_kwh > 1e-8)).sum())
    add("no_simultaneous_charge_discharge", simultaneous == 0, simultaneous)
    add("causal_surplus_deficit_rule", np.allclose(merged.emergency_purchase_kwh, np.maximum(merged.load_kw * DT - merged.grid_purchase_kwh - merged.pv_actual_kw * DT - merged.discharge_kwh, 0), atol=1e-6) and np.allclose(merged.spill_kwh, np.maximum(merged.grid_purchase_kwh + merged.pv_actual_kw * DT + merged.discharge_kwh - merged.load_kw * DT - merged.charge_kwh, 0), atol=1e-6))
    planned_cost = float(merged.grid_cost_yuan.sum())
    emergency_cost = float(merged.emergency_cost_yuan.sum())
    total_cost = planned_cost + emergency_cost
    add("cost_recalculation", np.isclose(emergency_cost, float((merged.emergency_purchase_kwh * 5 * merged.price_yuan_per_kwh).sum())) and np.isclose(total_cost, float(daily.total_cost_yuan.sum())))
    baseline = comparison.loc[comparison.strategy == "no_storage_same_forecast"].iloc[0]
    optimized = comparison.loc[comparison.strategy == "optimized_proposed"].iloc[0]
    raw = OUT / "result2_final.xlsx"
    wb = load_workbook(raw, read_only=True, data_only=False)
    row = [wb["计划购电量"].cell(2, c).value for c in range(2, 146)]
    first_plan = plan.loc[plan.date == START].sort_values("time_index").grid_purchase_kwh.to_numpy(float)
    raw_is_direct = np.allclose(row, first_plan)
    checked, cyclic = write_checked_workbook(plan, emergency)
    checked_wb = load_workbook(checked, read_only=True, data_only=False)
    checked_row = [checked_wb["计划购电量"].cell(2, c).value for c in range(2, 146)]
    add("excel_cyclic_headers_detected", cyclic)
    add("excel_checked_mapping", np.allclose(checked_row, first_plan[list(range(1, N)) + [0]]))
    metrics = pd.DataFrame([{
        "safety_quantile": 0.80, "feb1_initial_soc_kwh": float(daily.soc_start_kwh.iloc[0]),
        "planned_purchase_kwh": float(plan.grid_purchase_kwh.sum()), "planned_cost_yuan": planned_cost,
        "emergency_purchase_kwh": float(merged.emergency_purchase_kwh.sum()), "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": total_cost, "emergency_days": int((daily.actual_emergency_kwh > 1e-8).sum()),
        "emergency_intervals": int((merged.emergency_purchase_kwh > 1e-8).sum()), "spill_kwh": float(merged.spill_kwh.sum()),
        "no_storage_planned_cost_yuan": float(baseline.planned_cost_yuan), "no_storage_emergency_cost_yuan": float(baseline.emergency_cost_yuan), "no_storage_total_cost_yuan": float(baseline.total_cost_yuan),
        "saving_yuan": float(baseline.total_cost_yuan - total_cost), "saving_pct": float(100 * (baseline.total_cost_yuan - total_cost) / baseline.total_cost_yuan),
        "actual_soc_min_kwh": float(merged.soc_end_kwh.min()), "actual_soc_max_kwh": float(merged.soc_end_kwh.max()),
        "soc_continuity_max_error_kwh": max_continuity, "soc_transition_max_error_kwh": max_transition,
        "energy_balance_max_error_kwh": max_balance, "simultaneous_charge_discharge_intervals": simultaneous,
    }])
    metrics.to_csv(OUT / "q2_dynamic_battery_metrics.csv", index=False, encoding="utf-8-sig")
    report = ["# Q2 dynamic battery validation", "", "## Result", ""]
    report += [f"- {'PASS' if ok else 'FAIL'} — {name}" + (f": `{value:.3e}`" if value is not None else "") for name, ok, value in checks]
    report += ["", "## Excel time mapping", "", f"The original `result2_final.xlsx` used direct t=0..143 placement: **{'detected' if raw_is_direct else 'not detected'}**.", "The official headers are cyclic (t=1..143,0); `result2_final_checked.xlsx` preserves the raw workbook and corrects only Sheet 1 placement and Sheet 3 emergency-period labels.", "", "No terminal reserve or CVaR metric is evaluated. Causality is evidenced by strictly prior forecast history and by the recorded one-step dispatch balance."]
    (OUT / "q2_dynamic_battery_validation.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report))
    if not all(ok for _, ok, _ in checks): raise SystemExit(1)


if __name__ == "__main__": main()
