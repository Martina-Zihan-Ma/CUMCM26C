#!/usr/bin/env python3
"""Independent integrity checks for the Question 2 workflow and workbook."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "processed"
OUT = ROOT / "outputs" / "question2"
N, DT, ETA_C, ETA_D = 144, 1 / 6, .90, .90


def check(condition: bool, label: str, results: list[tuple[str, str]]) -> None:
    results.append(("PASS" if condition else "FAIL", label))


def complete_day_grid(frame: pd.DataFrame, date_col: str = "date") -> bool:
    counts = frame.groupby(date_col).size()
    indices = frame.groupby(date_col)["time_index"].apply(
        lambda x: sorted(x.astype(int).tolist()) == list(range(N))
    )
    return bool((counts == N).all() and indices.all())


def main() -> None:
    results: list[tuple[str, str]] = []
    price = pd.read_csv(DATA / "attachment1_standard_day.csv")
    actual = pd.read_csv(DATA / "attachment2_actual_long.csv")
    forecast = pd.read_csv(OUT / "dynamic_forecasts.csv")
    plan = pd.read_csv(OUT / "question2_plan.csv")
    dispatch = pd.read_csv(OUT / "question2_dispatch.csv")
    emergency = pd.read_csv(OUT / "question2_emergency.csv")
    daily = pd.read_csv(OUT / "question2_daily_summary.csv")

    for frame in (actual, forecast, plan, dispatch, emergency, daily):
        if "date" in frame:
            frame["date"] = pd.to_datetime(frame["date"]).dt.normalize()

    check(len(price) == N, "附件1恰为144行", results)
    check(price["time_index"].astype(int).tolist() == list(range(N)), "附件1 time_index 为0--143", results)
    check(len(actual) == 365 * N, "附件2恰为52560行", results)
    expected_dates = pd.date_range("2025-01-01", "2025-12-31", freq="D")
    check(actual["date"].drop_duplicates().tolist() == list(expected_dates), "附件2完整覆盖2025年", results)
    check(complete_day_grid(actual), "附件2每日144个完整且不重复时段", results)
    check(not actual.duplicated(["date", "time_index"]).any(), "附件2无重复日期-时段", results)
    check((actual[["load_kw", "pv_actual_kw"]] >= 0).all().all(), "附件2负载和光伏非负", results)
    check((price["price_yuan_per_kwh"] >= 0).all(), "电价非负", results)
    check(np.allclose(actual["load_kwh"], actual["load_kw"] * DT), "负载功率仅乘DT一次", results)
    check(np.allclose(actual["pv_actual_kwh"], actual["pv_actual_kw"] * DT), "光伏功率仅乘DT一次", results)

    check(len(forecast) == 364 * N and complete_day_grid(forecast), "预测输出覆盖Jan-2至Dec-31完整144时段", results)
    check((forecast["history_end_date"] < forecast["date"]).all(), "全部预测严格满足 history_end_date < target date", results)
    check((forecast[["forecast_load_kw", "forecast_pv_kw"]] >= 0).all().all(), "预测功率非负", results)
    check(np.allclose(forecast["load_residual_kw"], forecast["actual_load_kw"] - forecast["forecast_load_kw"]), "负载残差定义一致", results)
    check(np.allclose(forecast["pv_residual_kw"], forecast["actual_pv_kw"] - forecast["forecast_pv_kw"]), "光伏残差定义一致", results)

    operating_dates = pd.date_range("2025-02-01", "2025-12-31", freq="D")
    check(len(plan) == len(operating_dates) * N and complete_day_grid(plan), "计划购电覆盖334日×144时段", results)
    check(len(dispatch) == len(plan) and complete_day_grid(dispatch), "充放电调度覆盖334日×144时段", results)
    check(len(emergency) == len(plan) and complete_day_grid(emergency), "紧急购电结算覆盖334日×144时段", results)
    check((plan["grid_purchase_kwh"] >= -1e-8).all(), "计划购电量非负", results)
    check((dispatch[["charge_kwh", "discharge_kwh"]] >= -1e-8).all().all(), "充放电量非负", results)
    check((dispatch[["charge_kwh", "discharge_kwh"]].sum(axis=1) <= 5000 * DT + 1e-6).all(), "充放电共享功率上限满足", results)
    check((dispatch[["soc_start_kwh", "soc_end_kwh"]].to_numpy() >= 1200 - 1e-6).all() and (dispatch[["soc_start_kwh", "soc_end_kwh"]].to_numpy() <= 10800 + 1e-6).all(), "SOC上下界满足", results)
    transition = dispatch["soc_end_kwh"] - dispatch["soc_start_kwh"] - ETA_C * dispatch["charge_kwh"] + dispatch["discharge_kwh"] / ETA_D
    check(float(np.abs(transition).max()) < 1e-6, "SOC逐时段转移满足", results)
    check(float(np.abs(dispatch["soc_end_kwh"].iloc[:-1].to_numpy() - dispatch["soc_start_kwh"].iloc[1:].to_numpy()).max()) < 1e-6, "跨日SOC连续", results)

    merged = plan.merge(dispatch, on=["date", "time_index"]).merge(emergency, on=["date", "time_index"])
    actual_op = actual.loc[actual["date"].between(operating_dates[0], operating_dates[-1]), ["date", "time_index", "load_kw", "pv_actual_kw"]]
    merged = merged.merge(actual_op, on=["date", "time_index"])
    balance = merged["grid_purchase_kwh"] + merged["pv_actual_kw"] * DT + merged["discharge_kwh"] - merged["load_kw"] * DT - merged["charge_kwh"]
    expected_emergency = np.maximum(-balance.to_numpy(), 0)
    expected_spill = np.maximum(balance.to_numpy(), 0)
    check(np.allclose(merged["emergency_purchase_kwh"], expected_emergency, atol=1e-6), "实际紧急购电结算公式满足", results)
    check(np.allclose(merged["spill_kwh"], expected_spill, atol=1e-6), "实际弃电结算公式满足", results)
    check(np.allclose(merged["emergency_cost_yuan"], 5 * merged["price_yuan_per_kwh"] * merged["emergency_purchase_kwh"], atol=1e-6), "紧急购电按5倍电价结算", results)
    check(len(daily) == 334 and np.allclose(daily["total_cost_yuan"], daily["planned_cost_yuan"] + daily["actual_emergency_cost_yuan"]), "每日总费用为计划费用加实际紧急费用", results)

    workbook = OUT / "result2_final.xlsx"
    wb = load_workbook(workbook, read_only=True, data_only=False)
    ws1, ws2, ws3 = (wb[name] for name in ["计划购电量", "充放电量", "紧急购电量"])
    check(wb.sheetnames == ["计划购电量", "充放电量", "紧急购电量"], "结果工作簿工作表结构正确", results)
    check(ws1.max_row == 335 and ws1.max_column == 147, "Sheet1保留334天×144列模板结构", results)
    check(ws2.max_row == 1 + 334 * 6 and ws2.max_column == 6, "Sheet2已展开为334天六时段", results)
    check(ws3.max_row >= 335 and ws3.max_column == 3, "Sheet3已逐日写入紧急购电记录", results)
    first_day_sheet = [ws1.cell(2, c).value for c in range(2, 146)]
    first_day_plan = plan.loc[plan["date"] == operating_dates[0], "grid_purchase_kwh"].to_numpy()
    check(np.allclose(first_day_sheet, first_day_plan), "Sheet1首日购电量与CSV逐时段一致", results)

    report = ["Q2 RESULT VALIDATION", "=" * 20]
    report.extend(f"[{status}] {label}" for status, label in results)
    report.append("")
    report.append(f"Checks: {len(results)}; failures: {sum(s == 'FAIL' for s, _ in results)}")
    (OUT / "question2_validation.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report))
    if any(status == "FAIL" for status, _ in results):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
