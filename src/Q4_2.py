#!/usr/bin/env python3
"""
Q4-2 — dynamic-price version of Question 2.

This script intentionally preserves the Q2 forecasting/scenario/SOC framework and
makes only two conceptually separate improvements:

STEP 1: dynamic_greedy
    - Replace the single standard-day price vector with Attachment 4's date-specific
      144-slot price vector p[d,t].
    - Keep the original Q2 causal execution rule: surplus -> charge; deficit ->
      discharge as much as possible; remainder -> 5x emergency purchase.

STEP 2: dynamic_rolling (FINAL)
    - Keep the SAME day-ahead stochastic LP and fixed planned grid purchase G_t.
    - Replace the greedy real-time battery rule with a causal stochastic
      receding-horizon battery dispatch.
    - At time t, only the CURRENT realized Load/PV and current SOC are known.
      Future Load/PV remain the 0:00 bootstrap scenarios; future actual values are
      never used.
    - Grid purchase G is never re-planned.  The rolling LP decides only current
      charge/discharge/emergency/spill by valuing battery energy against expected
      future 5x emergency-purchase cost under the dynamic price curve.
    - No hand-built reserve path and no future-max heuristic are used.

No CVaR term is added here.  The Q2 risk structure is preserved:
    planned grid cost + expected 5x emergency cost
plus the existing SAFETY_QUANTILE hard safety-scenario constraint.

Expected project layout
-----------------------
project/
  data/processed/attachment2_actual_long.csv
  data/processed/attachment4_price_long.csv
  outputs/question2/dynamic_forecasts.csv
  results/result4-2_template.xlsx   (preferred name)
  src/<this file>

The template finder also accepts result4-2.xlsx and result4-2(1).xlsx.
"""

from __future__ import annotations

import os
import time
from copy import copy, deepcopy
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


# ======================================================================
# Paths / calendar
# ======================================================================
ROOT = Path(__file__).resolve().parents[1]
PRICE_FILE = ROOT / "data" / "processed" / "attachment4_price_long.csv"
ACTUAL_FILE = ROOT / "data" / "processed" / "attachment2_actual_long.csv"
FORECAST_FILE = ROOT / "outputs" / "question2" / "dynamic_forecasts.csv"
OUT = ROOT / "outputs" / "question4_2"

JAN1 = pd.Timestamp("2025-01-01")
WARMUP_START = pd.Timestamp("2025-01-02")
WARMUP_END = pd.Timestamp("2025-01-31")
START = pd.Timestamp("2025-02-01")
END = pd.Timestamp("2025-12-31")
DATES = pd.date_range(START, END, freq="D")

N = 144
DT = 1.0 / 6.0  # 10 minutes = 1/6 hour


# ======================================================================
# Battery / market parameters (UNCHANGED from Q2)
# ======================================================================
ETA_C = 0.90
ETA_D = 0.90
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_JAN1 = 6000.0
MAX_INTERVAL = 5000.0 * DT  # kWh in one 10-min interval

EMERGENCY_MULTIPLIER = 5.0


# ======================================================================
# Q2 uncertainty structure (UNCHANGED)
# ======================================================================
RESIDUAL_DAYS = 30
RESIDUAL_HALF_LIFE = 14.0
SCENARIOS = 20
MIN_RESIDUAL_DAYS = 5
RNG_SEED = 20260911

# Existing Q2 hard safety rule: the sampled whole-day scenario at this
# severity quantile must be feasible with zero emergency purchase.
SAFETY_QUANTILE = 0.80

# Tiny LP tie-breaking penalties.
CYCLE_EPS = 1e-6
SPILL_EPS = 1e-8

# Debug: Q4_2_MAX_DAYS=5 python src/Q4_2_dynamic_price.py
MAX_DAYS = int(os.environ.get("Q4_2_MAX_DAYS", "0"))
PROGRESS_EVERY_DAYS = max(1, int(os.environ.get("Q4_PROGRESS_EVERY", "1")))

# Final official workbook uses the rolling stochastic execution strategy.
FINAL_STRATEGY_NAME = "dynamic_rolling"

# Rolling controller horizon. 0 means use the whole remaining day.
# A finite cap can be supplied for speed, e.g.
#   Q4_ROLLING_HORIZON=48 python src/Q4_2.py
# The default uses the full remaining day for the cleanest economic comparison.
ROLLING_HORIZON = int(os.environ.get("Q4_ROLLING_HORIZON", "0"))


# ======================================================================
# Input
# ======================================================================
def read_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    for path in [PRICE_FILE, ACTUAL_FILE, FORECAST_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    price = pd.read_csv(PRICE_FILE)
    actual = pd.read_csv(ACTUAL_FILE)
    forecast = pd.read_csv(FORECAST_FILE)

    price["date"] = pd.to_datetime(price["date"]).dt.normalize()
    actual["date"] = pd.to_datetime(actual["date"]).dt.normalize()
    forecast["date"] = pd.to_datetime(forecast["date"]).dt.normalize()

    if "history_end_date" in forecast.columns:
        forecast["history_end_date"] = pd.to_datetime(
            forecast["history_end_date"]
        ).dt.normalize()

    required_price = {"date", "time_index", "price_yuan_per_kwh"}
    required_actual = {"date", "time_index", "load_kw", "pv_actual_kw"}
    required_forecast = {
        "date", "time_index", "forecast_load_kw", "forecast_pv_kw",
        "load_residual_kw", "pv_residual_kw",
    }

    for name, df, required in [
        ("price", price, required_price),
        ("actual", actual, required_actual),
        ("forecast", forecast, required_forecast),
    ]:
        missing = required - set(df.columns)
        if missing:
            raise ValueError(f"{name} missing columns: {sorted(missing)}")

    # Attachment 4 must contain exactly one 144-slot price vector per day.
    expected_dates = pd.date_range(JAN1, END, freq="D")
    present_dates = pd.DatetimeIndex(sorted(price["date"].unique()))
    if not expected_dates.equals(present_dates):
        raise ValueError("Attachment 4 price file must contain every day of 2025.")

    for date, part in price.groupby("date"):
        idx = part.sort_values("time_index")["time_index"].astype(int).tolist()
        if idx != list(range(N)):
            raise ValueError(f"price {date.date()}: time_index must be 0..143 exactly once.")

    for name, df in [("actual", actual), ("forecast", forecast)]:
        for date, part in df.groupby("date"):
            idx = part.sort_values("time_index")["time_index"].astype(int).tolist()
            if idx != list(range(N)):
                raise ValueError(f"{name} {date.date()}: time_index must be 0..143.")

    if forecast["date"].min() > WARMUP_START:
        raise ValueError("Forecast file must include January warm-up forecasts from Jan 2.")
    if forecast["date"].max() < END:
        raise ValueError("Forecast file must include forecasts through Dec 31.")

    if "history_end_date" in forecast.columns:
        if not (forecast["history_end_date"] < forecast["date"]).all():
            raise RuntimeError("Forecast leakage: history_end_date >= target date.")

    return price, actual, forecast


def day_values(df: pd.DataFrame, date: pd.Timestamp, col: str) -> np.ndarray:
    x = df.loc[df["date"] == date].sort_values("time_index")
    if len(x) != N:
        raise ValueError(f"{date.date()}: {col} has {len(x)} rows, expected {N}.")
    return x[col].to_numpy(float)


def day_price(price: pd.DataFrame, date: pd.Timestamp) -> np.ndarray:
    p = day_values(price, date, "price_yuan_per_kwh")
    if np.any(~np.isfinite(p)) or np.any(p < 0):
        raise ValueError(f"{date.date()}: invalid price values.")
    return p


def proposed_forecast(
    forecast: pd.DataFrame,
    date: pd.Timestamp,
) -> tuple[np.ndarray, np.ndarray]:
    return (
        day_values(forecast, date, "forecast_load_kw"),
        day_values(forecast, date, "forecast_pv_kw"),
    )


# ======================================================================
# Historical forecast residuals (UNCHANGED)
# ======================================================================
def positive_net_energy(load_kw: np.ndarray, pv_kw: np.ndarray) -> float:
    return float(np.maximum(load_kw - pv_kw, 0.0).sum() * DT)


def forecast_error_risk(
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
) -> float:
    """Daily positive actual-minus-forecast net-demand error (kWh)."""
    net_error_kw = (load_actual - load_hat) - (pv_actual - pv_hat)
    return float(np.maximum(net_error_kw, 0.0).sum() * DT)


def residual_record(
    date: pd.Timestamp,
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
) -> dict:
    return {
        "date": date.normalize(),
        "load_error_kw": (load_actual - load_hat).astype(float),
        "pv_error_kw": (pv_actual - pv_hat).astype(float),
        "forecast_net_kwh": positive_net_energy(load_hat, pv_hat),
        "risk_kwh": forecast_error_risk(
            load_hat, pv_hat, load_actual, pv_actual
        ),
    }


def recent_residuals(history: list[dict], date: pd.Timestamp) -> list[dict]:
    prior = sorted(
        [r for r in history if r["date"] < date],
        key=lambda r: r["date"],
    )
    return prior[-RESIDUAL_DAYS:]


# ======================================================================
# Residual bootstrap scenarios (UNCHANGED)
# ======================================================================
def sample_scenarios(
    date: pd.Timestamp,
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    residual_history: list[dict],
    rng: np.random.Generator,
) -> tuple[np.ndarray, np.ndarray]:
    residuals = recent_residuals(residual_history, date)

    if len(residuals) < MIN_RESIDUAL_DAYS:
        return (
            np.maximum(load_hat, 0.0)[None, :] * DT,
            np.maximum(pv_hat, 0.0)[None, :] * DT,
        )

    ages = np.array([(date - r["date"]).days for r in residuals], dtype=float)
    weights = np.exp(-np.log(2.0) * ages / RESIDUAL_HALF_LIFE)
    weights /= weights.sum()

    selected = rng.choice(
        len(residuals), size=SCENARIOS, replace=True, p=weights
    )

    load_scenarios = []
    pv_scenarios = []
    for i in selected:
        r = residuals[int(i)]
        load_scenarios.append(
            np.maximum(load_hat + r["load_error_kw"], 0.0) * DT
        )
        pv_scenarios.append(
            np.maximum(pv_hat + r["pv_error_kw"], 0.0) * DT
        )

    return np.vstack(load_scenarios), np.vstack(pv_scenarios)


# ======================================================================
# Day-ahead stochastic LP
# ======================================================================
@dataclass
class VarIndex:
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    soc: np.ndarray
    emergency: np.ndarray
    spill: np.ndarray


def make_index(k: int) -> tuple[VarIndex, int]:
    p = 0
    grid = np.arange(p, p + N); p += N
    charge = np.arange(p, p + k * N).reshape(k, N); p += k * N
    discharge = np.arange(p, p + k * N).reshape(k, N); p += k * N
    soc = np.arange(p, p + k * (N + 1)).reshape(k, N + 1); p += k * (N + 1)
    emergency = np.arange(p, p + k * N).reshape(k, N); p += k * N
    spill = np.arange(p, p + k * N).reshape(k, N); p += k * N
    return VarIndex(grid, charge, discharge, soc, emergency, spill), p


def choose_safety_scenario(net_scenarios: np.ndarray) -> tuple[int, float]:
    """Choose the sampled whole-day scenario at SAFETY_QUANTILE severity."""
    severity = np.maximum(net_scenarios, 0.0).sum(axis=1)
    order = np.argsort(severity)
    rank = int(np.ceil(SAFETY_QUANTILE * len(order))) - 1
    rank = int(np.clip(rank, 0, len(order) - 1))
    idx = int(order[rank])
    return idx, float(severity[idx])


def solve_plan(
    price: np.ndarray,
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    initial_soc: float,
) -> dict:
    """
    Same Q2 day-ahead stochastic LP, now using the DATE-SPECIFIC price vector.

    First-stage decision:
        grid[t] fixed at 0:00.

    Scenario recourse:
        charge/discharge/SOC/emergency/spill.

    Objective (NO CVaR):
        planned grid cost + expected 5x emergency cost.

    The sampled load/PV scenarios are also returned so Step 2 can reuse exactly
    the same 0:00 uncertainty information in its causal rolling battery controller.
    """
    k = load_scenarios.shape[0]
    if load_scenarios.shape != (k, N) or pv_scenarios.shape != (k, N):
        raise ValueError("Bad scenario shape.")

    net_scenarios = load_scenarios - pv_scenarios
    safety_idx, safety_energy = choose_safety_scenario(net_scenarios)

    ix, nv = make_index(k)

    c = np.zeros(nv)
    c[ix.grid] = price
    for w in range(k):
        c[ix.charge[w]] = CYCLE_EPS / k
        c[ix.discharge[w]] = CYCLE_EPS / k
        c[ix.emergency[w]] = EMERGENCY_MULTIPLIER * price / k
        c[ix.spill[w]] = SPILL_EPS / k

    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)

    for w in range(k):
        ub[ix.charge[w]] = MAX_INTERVAL
        ub[ix.discharge[w]] = MAX_INTERVAL
        lb[ix.soc[w]] = SOC_MIN
        ub[ix.soc[w]] = SOC_MAX
        lb[ix.soc[w, 0]] = initial_soc
        ub[ix.soc[w, 0]] = initial_soc

    # Existing Q2 safety protection: selected safety scenario has no emergency.
    ub[ix.emergency[safety_idx]] = 0.0

    # Energy balances + SOC dynamics.
    Aeq = lil_matrix((2 * k * N, nv), dtype=float)
    beq = np.zeros(2 * k * N, dtype=float)
    row = 0

    for w in range(k):
        for t in range(N):
            # G + E + D = (Load - PV) + C + Spill
            Aeq[row, ix.grid[t]] = 1.0
            Aeq[row, ix.emergency[w, t]] = 1.0
            Aeq[row, ix.discharge[w, t]] = 1.0
            Aeq[row, ix.charge[w, t]] = -1.0
            Aeq[row, ix.spill[w, t]] = -1.0
            beq[row] = net_scenarios[w, t]
            row += 1

    for w in range(k):
        for t in range(N):
            # SOC_{t+1} = SOC_t + eta_c*C_t - D_t/eta_d
            Aeq[row, ix.soc[w, t + 1]] = 1.0
            Aeq[row, ix.soc[w, t]] = -1.0
            Aeq[row, ix.charge[w, t]] = -ETA_C
            Aeq[row, ix.discharge[w, t]] = 1.0 / ETA_D
            row += 1

    # Shared charge/discharge power envelope per scenario/slot.
    Aub = lil_matrix((k * N, nv), dtype=float)
    bub = np.full(k * N, MAX_INTERVAL, dtype=float)
    row = 0
    for w in range(k):
        for t in range(N):
            Aub[row, ix.charge[w, t]] = 1.0
            Aub[row, ix.discharge[w, t]] = 1.0
            row += 1

    result = linprog(
        c=c,
        A_ub=Aub.tocsr(),
        b_ub=bub,
        A_eq=Aeq.tocsr(),
        b_eq=beq,
        bounds=list(zip(lb, ub)),
        method="highs",
    )
    if not result.success:
        raise RuntimeError(result.message)

    x = result.x
    grid = x[ix.grid].copy()

    scenario_emergency_cost = np.array([
        np.dot(x[ix.emergency[w]], EMERGENCY_MULTIPLIER * price)
        for w in range(k)
    ])

    return {
        "grid": grid,
        "expected_emergency_cost": float(scenario_emergency_cost.mean()),
        "safety_scenario_index": safety_idx,
        "safety_scenario_positive_net_kwh": safety_energy,
        "safety_scenario_emergency_kwh": float(x[ix.emergency[safety_idx]].sum()),
        "load_scenarios": load_scenarios.copy(),
        "pv_scenarios": pv_scenarios.copy(),
    }


# ======================================================================
# STEP 1 execution: original Q2 greedy causal rule
# ======================================================================
def execute_dynamic_greedy(
    grid: np.ndarray,
    initial_soc: float,
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    plan: dict | None = None,
) -> dict:
    """Original Q2 execution; dynamic price affects cost but not battery rule."""
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    spill = np.zeros(N)
    soc_start = np.zeros(N)
    soc_end = np.zeros(N)
    rolling_objective = np.full(N, np.nan)
    rolling_horizon = np.zeros(N, dtype=int)

    soc = float(initial_soc)

    for t in range(N):
        soc_start[t] = soc
        balance = float(grid[t] + pv_kwh[t] - load_kwh[t])

        if balance >= 0.0:
            capacity_input = max((SOC_MAX - soc) / ETA_C, 0.0)
            charge[t] = min(balance, MAX_INTERVAL, capacity_input)
            soc += ETA_C * charge[t]
            spill[t] = max(balance - charge[t], 0.0)
        else:
            deficit = -balance
            available_output = max((soc - SOC_MIN) * ETA_D, 0.0)
            discharge[t] = min(deficit, MAX_INTERVAL, available_output)
            soc -= discharge[t] / ETA_D
            emergency[t] = max(deficit - discharge[t], 0.0)

        soc = float(np.clip(soc, SOC_MIN, SOC_MAX))
        soc_end[t] = soc

    balance_error = np.max(np.abs(
        grid + emergency + pv_kwh + discharge
        - load_kwh - charge - spill
    ))

    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "soc_start": soc_start,
        "soc_end": soc_end,
        "terminal_soc": float(soc),
        "balance_error": float(balance_error),
        "rolling_objective": rolling_objective,
        "rolling_horizon": rolling_horizon,
    }


# ======================================================================
# STEP 2 execution: stochastic receding-horizon battery control
# ======================================================================
def _rolling_first_action(
    t: int,
    grid: np.ndarray,
    price: np.ndarray,
    current_soc: float,
    current_net_kwh: float,
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
) -> tuple[float, float, float, float, float, int]:
    """
    Solve a stochastic look-ahead LP and return ONLY the action for slot t.

    Information structure
    ---------------------
    - current_net_kwh is the realized net demand at t and is identical in every
      scenario;
    - future net demand uses only the scenarios sampled at 0:00;
    - grid[t:] is fixed to the day-ahead plan and is never optimized here;
    - scenario-specific future recourse is allowed, but the first action
      (charge/discharge/emergency/spill at t) is constrained to be identical
      across scenarios (one-step non-anticipativity).

    Objective
    ---------
    Minimize expected 5x emergency-purchase cost over the remaining horizon.
    Planned grid cost is already sunk/fixed, so it does not enter this LP.
    """
    k = int(load_scenarios.shape[0])
    if load_scenarios.shape != pv_scenarios.shape or load_scenarios.shape[1] != N:
        raise ValueError("Bad rolling scenario shape.")

    if ROLLING_HORIZON > 0:
        end = min(N, t + ROLLING_HORIZON)
    else:
        end = N
    h = end - t
    if h <= 0:
        raise ValueError("Empty rolling horizon.")

    # Scenario net demand over the look-ahead window (kWh/slot).
    net = load_scenarios[:, t:end] - pv_scenarios[:, t:end]
    net = np.asarray(net, dtype=float).copy()
    net[:, 0] = float(current_net_kwh)  # current actual replaces forecast/scenario

    # Variables per scenario and horizon slot.
    # C, D, E, Spill: (k,h); SOC: (k,h+1)
    p = 0
    charge = np.arange(p, p + k * h).reshape(k, h); p += k * h
    discharge = np.arange(p, p + k * h).reshape(k, h); p += k * h
    emergency = np.arange(p, p + k * h).reshape(k, h); p += k * h
    spill = np.arange(p, p + k * h).reshape(k, h); p += k * h
    soc = np.arange(p, p + k * (h + 1)).reshape(k, h + 1); p += k * (h + 1)
    nv = p

    c = np.zeros(nv)
    p_window = price[t:end]
    for w in range(k):
        c[charge[w]] = CYCLE_EPS / k
        c[discharge[w]] = CYCLE_EPS / k
        c[emergency[w]] = EMERGENCY_MULTIPLIER * p_window / k
        c[spill[w]] = SPILL_EPS / k

    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)
    for w in range(k):
        ub[charge[w]] = MAX_INTERVAL
        ub[discharge[w]] = MAX_INTERVAL
        lb[soc[w]] = SOC_MIN
        ub[soc[w]] = SOC_MAX
        lb[soc[w, 0]] = current_soc
        ub[soc[w, 0]] = current_soc

    # Equality constraints:
    # 1) energy balance in every scenario/slot
    # 2) SOC dynamics in every scenario/slot
    # 3) first action is common across scenarios (non-anticipativity)
    nonant_vars = [charge, discharge, emergency, spill]
    n_nonant = max(k - 1, 0) * len(nonant_vars)
    Aeq = lil_matrix((2 * k * h + n_nonant, nv), dtype=float)
    beq = np.zeros(2 * k * h + n_nonant, dtype=float)
    row = 0

    for w in range(k):
        for j in range(h):
            # E + D - C - Spill = Net - fixed Grid
            Aeq[row, emergency[w, j]] = 1.0
            Aeq[row, discharge[w, j]] = 1.0
            Aeq[row, charge[w, j]] = -1.0
            Aeq[row, spill[w, j]] = -1.0
            beq[row] = net[w, j] - grid[t + j]
            row += 1

    for w in range(k):
        for j in range(h):
            Aeq[row, soc[w, j + 1]] = 1.0
            Aeq[row, soc[w, j]] = -1.0
            Aeq[row, charge[w, j]] = -ETA_C
            Aeq[row, discharge[w, j]] = 1.0 / ETA_D
            row += 1

    for w in range(1, k):
        for arr in nonant_vars:
            Aeq[row, arr[w, 0]] = 1.0
            Aeq[row, arr[0, 0]] = -1.0
            row += 1

    # Charge + discharge cannot exceed inverter power in a slot.
    Aub = lil_matrix((k * h, nv), dtype=float)
    bub = np.full(k * h, MAX_INTERVAL, dtype=float)
    row = 0
    for w in range(k):
        for j in range(h):
            Aub[row, charge[w, j]] = 1.0
            Aub[row, discharge[w, j]] = 1.0
            row += 1

    result = linprog(
        c=c,
        A_ub=Aub.tocsr(),
        b_ub=bub,
        A_eq=Aeq.tocsr(),
        b_eq=beq,
        bounds=list(zip(lb, ub)),
        method="highs",
    )
    if not result.success:
        raise RuntimeError(
            f"Rolling LP failed at slot {t}: {result.message}"
        )

    x = result.x
    # Because of non-anticipativity, scenario 0's first action equals all others.
    return (
        float(x[charge[0, 0]]),
        float(x[discharge[0, 0]]),
        float(x[emergency[0, 0]]),
        float(x[spill[0, 0]]),
        float(result.fun),
        int(h),
    )


def execute_rolling_stochastic(
    grid: np.ndarray,
    initial_soc: float,
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    plan: dict,
) -> dict:
    """
    Strictly causal stochastic receding-horizon battery dispatch.

    The day-ahead purchase vector grid[] is NEVER changed.  At each 10-minute
    slot, the current actual Load/PV and SOC are observed, then a stochastic
    look-ahead LP decides only the current battery/emergency/spill action.  Future
    actual Load/PV are not accessed; the future remains the 0:00 scenarios.

    To avoid unnecessary LP solves, obvious cases are handled directly:
      - current surplus: charge greedily (free sunk surplus cannot be sold);
      - no usable battery during a deficit: emergency purchase is forced;
      - otherwise solve the rolling LP.
    """
    load_scenarios = np.asarray(plan["load_scenarios"], dtype=float)
    pv_scenarios = np.asarray(plan["pv_scenarios"], dtype=float)

    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    spill = np.zeros(N)
    soc_start = np.zeros(N)
    soc_end = np.zeros(N)
    rolling_objective = np.full(N, np.nan)
    rolling_horizon = np.zeros(N, dtype=int)

    soc_now = float(initial_soc)

    for t in range(N):
        soc_start[t] = soc_now
        current_net = float(load_kwh[t] - pv_kwh[t])
        balance = float(grid[t] - current_net)  # grid + PV - Load

        if balance >= -1e-12:
            # Surplus grid/PV is already available and cannot be sold.  Charging
            # it is weakly better than spilling while storage has room.
            capacity_input = max((SOC_MAX - soc_now) / ETA_C, 0.0)
            charge[t] = min(max(balance, 0.0), MAX_INTERVAL, capacity_input)
            soc_now += ETA_C * charge[t]
            spill[t] = max(balance - charge[t], 0.0)
        else:
            deficit = -balance
            physical_output = max((soc_now - SOC_MIN) * ETA_D, 0.0)

            if physical_output <= 1e-12:
                # No battery energy can legally be discharged.
                emergency[t] = deficit
            else:
                c0, d0, e0, s0, obj, h = _rolling_first_action(
                    t=t,
                    grid=grid,
                    price=price,
                    current_soc=soc_now,
                    current_net_kwh=current_net,
                    load_scenarios=load_scenarios,
                    pv_scenarios=pv_scenarios,
                )

                # Numerical cleanup; the LP itself enforces all physical bounds.
                charge[t] = max(c0, 0.0)
                discharge[t] = max(d0, 0.0)
                emergency[t] = max(e0, 0.0)
                spill[t] = max(s0, 0.0)
                rolling_objective[t] = obj
                rolling_horizon[t] = h

                soc_now += ETA_C * charge[t] - discharge[t] / ETA_D

        soc_now = float(np.clip(soc_now, SOC_MIN, SOC_MAX))
        soc_end[t] = soc_now

    balance_error = np.max(np.abs(
        grid + emergency + pv_kwh + discharge
        - load_kwh - charge - spill
    ))

    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "soc_start": soc_start,
        "soc_end": soc_end,
        "terminal_soc": float(soc_now),
        "balance_error": float(balance_error),
        "rolling_objective": rolling_objective,
        "rolling_horizon": rolling_horizon,
    }


# ======================================================================
# January warm-up
# ======================================================================
def warm_up_january(
    price_df: pd.DataFrame,
    actual: pd.DataFrame,
    forecast: pd.DataFrame,
    execution_mode: str,
) -> tuple[float, list[dict], pd.DataFrame]:
    rng = np.random.default_rng(RNG_SEED)
    soc = SOC_JAN1
    residuals: list[dict] = []

    rows = [{
        "date": JAN1,
        "soc_start_kwh": SOC_JAN1,
        "soc_end_kwh": SOC_JAN1,
        "note": "initialization day; no prior history",
    }]

    executor = (
        execute_dynamic_greedy
        if execution_mode == "dynamic_greedy"
        else execute_rolling_stochastic
    )

    warmup_dates = list(pd.date_range(WARMUP_START, WARMUP_END, freq="D"))
    for day_i, date in enumerate(warmup_dates, start=1):
        day_t0 = time.perf_counter()
        if execution_mode == "dynamic_rolling":
            print(
                f"[warm-up {day_i:02d}/{len(warmup_dates)}] {date.date()} "
                f"start  SOC0={soc:,.1f} kWh",
                flush=True,
            )

        p = day_price(price_df, date)
        load_hat, pv_hat = proposed_forecast(forecast, date)
        load_s, pv_s = sample_scenarios(date, load_hat, pv_hat, residuals, rng)
        plan = solve_plan(p, load_s, pv_s, soc)

        load_actual_kw = day_values(actual, date, "load_kw")
        pv_actual_kw = day_values(actual, date, "pv_actual_kw")
        dispatch = executor(
            plan["grid"], soc,
            load_actual_kw * DT,
            pv_actual_kw * DT,
            p,
            plan,
        )

        rows.append({
            "date": date,
            "soc_start_kwh": soc,
            "soc_end_kwh": dispatch["terminal_soc"],
            "planned_purchase_kwh": float(plan["grid"].sum()),
            "actual_emergency_kwh": float(dispatch["emergency"].sum()),
            "note": execution_mode,
        })

        # Only AFTER this day finishes may its actual error enter later scenarios.
        residuals.append(residual_record(
            date, load_hat, pv_hat, load_actual_kw, pv_actual_kw
        ))
        soc = dispatch["terminal_soc"]

        if execution_mode == "dynamic_rolling":
            lp_calls = int(np.isfinite(dispatch["rolling_objective"]).sum())
            elapsed = time.perf_counter() - day_t0
            print(
                f"[warm-up {day_i:02d}/{len(warmup_dates)}] {date.date()} "
                f"done   LP={lp_calls:3d}  emergency={dispatch['emergency'].sum():,.1f} kWh  "
                f"SOC24={soc:,.1f}  elapsed={elapsed:.1f}s",
                flush=True,
            )

    return float(soc), residuals, pd.DataFrame(rows)


# ======================================================================
# Feb-Dec strategy runner
# ======================================================================
def run_strategy(
    name: str,
    price_df: pd.DataFrame,
    actual: pd.DataFrame,
    forecast: pd.DataFrame,
    initial_soc: float,
    january_residuals: list[dict],
    dates: list[pd.Timestamp],
    keep_detail: bool,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if name not in {"dynamic_greedy", "dynamic_rolling"}:
        raise ValueError(name)

    executor = (
        execute_dynamic_greedy
        if name == "dynamic_greedy"
        else execute_rolling_stochastic
    )

    rng = np.random.default_rng(RNG_SEED + 1)
    residuals = deepcopy(january_residuals)
    soc = float(initial_soc)

    plan_rows: list[dict] = []
    dispatch_rows: list[dict] = []
    emergency_rows: list[dict] = []
    rolling_rows: list[dict] = []
    daily_rows: list[dict] = []
    max_balance_error = 0.0

    total_days = len(dates)
    strategy_t0 = time.perf_counter()
    for day_i, date in enumerate(dates, start=1):
        day_t0 = time.perf_counter()
        if day_i == 1 or day_i % PROGRESS_EVERY_DAYS == 0 or day_i == total_days:
            print(
                f"[{name} {day_i:03d}/{total_days}] {date.date()} start  "
                f"SOC0={soc:,.1f} kWh",
                flush=True,
            )

        p = day_price(price_df, date)
        load_hat, pv_hat = proposed_forecast(forecast, date)
        load_s, pv_s = sample_scenarios(date, load_hat, pv_hat, residuals, rng)

        # Binding Q4-2 day-ahead planned grid purchase, fixed at 0:00.
        plan = solve_plan(p, load_s, pv_s, soc)

        # Strictly causal execution: current actual only, no future actual leakage.
        load_actual_kw = day_values(actual, date, "load_kw")
        pv_actual_kw = day_values(actual, date, "pv_actual_kw")
        dispatch = executor(
            plan["grid"], soc,
            load_actual_kw * DT,
            pv_actual_kw * DT,
            p,
            plan,
        )

        emergency = dispatch["emergency"]
        spill = dispatch["spill"]
        max_balance_error = max(max_balance_error, dispatch["balance_error"])

        planned_cost = float(np.dot(plan["grid"], p))
        emergency_cost = float(np.dot(
            emergency, EMERGENCY_MULTIPLIER * p
        ))

        daily_rows.append({
            "strategy": name,
            "date": date,
            "soc_start_kwh": soc,
            "soc_end_kwh": dispatch["terminal_soc"],
            "planned_purchase_kwh": float(plan["grid"].sum()),
            "planned_cost_yuan": planned_cost,
            "expected_emergency_cost_yuan": plan["expected_emergency_cost"],
            "safety_quantile": SAFETY_QUANTILE,
            "safety_scenario_positive_net_kwh": plan[
                "safety_scenario_positive_net_kwh"
            ],
            "actual_emergency_kwh": float(emergency.sum()),
            "actual_emergency_cost_yuan": emergency_cost,
            "actual_emergency_intervals": int((emergency > 1e-8).sum()),
            "spill_kwh": float(spill.sum()),
            "rolling_lp_calls": int(np.isfinite(dispatch["rolling_objective"]).sum()),
            "total_cost_yuan": planned_cost + emergency_cost,
            "day_price_min": float(p.min()),
            "day_price_max": float(p.max()),
            "day_price_mean": float(p.mean()),
        })

        if keep_detail:
            for t in range(N):
                plan_rows.append({
                    "date": date,
                    "time_index": t,
                    "grid_purchase_kwh": float(plan["grid"][t]),
                    "price_yuan_per_kwh": float(p[t]),
                    "grid_cost_yuan": float(plan["grid"][t] * p[t]),
                })

                dispatch_rows.append({
                    "date": date,
                    "time_index": t,
                    "charge_kwh": float(dispatch["charge"][t]),
                    "discharge_kwh": float(dispatch["discharge"][t]),
                    "soc_start_kwh": float(dispatch["soc_start"][t]),
                    "soc_end_kwh": float(dispatch["soc_end"][t]),
                    "spill_kwh": float(spill[t]),
                    "rolling_objective_yuan": (
                        float(dispatch["rolling_objective"][t])
                        if np.isfinite(dispatch["rolling_objective"][t]) else np.nan
                    ),
                    "rolling_horizon_slots": int(dispatch["rolling_horizon"][t]),
                })

                emergency_rows.append({
                    "date": date,
                    "time_index": t,
                    "emergency_purchase_kwh": float(emergency[t]),
                    "price_yuan_per_kwh": float(p[t]),
                    "emergency_cost_yuan": float(
                        emergency[t] * EMERGENCY_MULTIPLIER * p[t]
                    ),
                })

                rolling_rows.append({
                    "date": date,
                    "time_index": t,
                    "rolling_lp_called": int(np.isfinite(dispatch["rolling_objective"][t])),
                    "rolling_objective_yuan": (
                        float(dispatch["rolling_objective"][t])
                        if np.isfinite(dispatch["rolling_objective"][t]) else np.nan
                    ),
                    "rolling_horizon_slots": int(dispatch["rolling_horizon"][t]),
                })

        residuals.append(residual_record(
            date, load_hat, pv_hat, load_actual_kw, pv_actual_kw
        ))
        soc = dispatch["terminal_soc"]

        if day_i == 1 or day_i % PROGRESS_EVERY_DAYS == 0 or day_i == total_days:
            elapsed = time.perf_counter() - day_t0
            total_elapsed = time.perf_counter() - strategy_t0
            lp_calls = int(np.isfinite(dispatch["rolling_objective"]).sum())
            print(
                f"[{name} {day_i:03d}/{total_days}] {date.date()} done   "
                f"LP={lp_calls:3d}  emergency={emergency.sum():,.1f} kWh  "
                f"SOC24={soc:,.1f}  day={elapsed:.1f}s  total={total_elapsed/60:.1f}min",
                flush=True,
            )

    daily = pd.DataFrame(daily_rows)
    daily.attrs["max_balance_error"] = max_balance_error

    return (
        pd.DataFrame(plan_rows),
        pd.DataFrame(dispatch_rows),
        pd.DataFrame(emergency_rows),
        pd.DataFrame(rolling_rows),
        daily,
    )


def summarize_strategy(daily: pd.DataFrame) -> dict:
    emergency_kwh = float(daily["actual_emergency_kwh"].sum())
    emergency_cost = float(daily["actual_emergency_cost_yuan"].sum())
    return {
        "strategy": daily["strategy"].iloc[0],
        "planned_purchase_kwh": float(daily["planned_purchase_kwh"].sum()),
        "planned_cost_yuan": float(daily["planned_cost_yuan"].sum()),
        "emergency_purchase_kwh": emergency_kwh,
        "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": float(daily["total_cost_yuan"].sum()),
        "spill_kwh": float(daily["spill_kwh"].sum()),
        "emergency_days": int((daily["actual_emergency_kwh"] > 1e-8).sum()),
        "emergency_intervals": int(daily["actual_emergency_intervals"].sum()),
        "rolling_lp_calls": int(daily["rolling_lp_calls"].sum()),
        "avg_normal_price_of_emergency": (
            emergency_cost / (EMERGENCY_MULTIPLIER * emergency_kwh)
            if emergency_kwh > 1e-12 else np.nan
        ),
    }


# ======================================================================
# Official workbook utilities
# ======================================================================
def _format_natural_time(total_minutes: int) -> str:
    total_minutes = int(total_minutes)
    if total_minutes == 24 * 60:
        return "24:00"
    hour, minute = divmod(total_minutes, 60)
    return f"{hour}:{minute:02d}"


def compress_emergency(emergency: pd.DataFrame) -> pd.DataFrame:
    rows = []
    active = emergency.loc[emergency["emergency_purchase_kwh"] > 1e-8].copy()
    if active.empty:
        return pd.DataFrame(
            columns=["date", "start_index", "end_index", "period", "emergency_kwh"]
        )

    for date, day in active.groupby("date"):
        day = day.sort_values("time_index")
        idx = day["time_index"].astype(int).tolist()
        values = day["emergency_purchase_kwh"].to_numpy(float)

        start = 0
        for j in range(1, len(idx) + 1):
            end_group = j == len(idx) or idx[j] != idx[j - 1] + 1
            if not end_group:
                continue

            a = idx[start]
            last = idx[j - 1]
            start_min = a * 10
            end_min = (last + 1) * 10

            rows.append({
                "date": pd.Timestamp(date).normalize(),
                "start_index": int(a),
                "end_index": int(last),
                "period": (
                    f"{_format_natural_time(start_min)}-"
                    f"{_format_natural_time(end_min)}"
                ),
                "emergency_kwh": float(values[start:j].sum()),
            })
            start = j

    return pd.DataFrame(rows)


def find_template() -> Path:
    candidates = [
        ROOT / "results" / "result4-2_template.xlsx",
        ROOT / "results" / "result4-2.xlsx",
        ROOT / "results" / "result4-2(1).xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(
        "Cannot find Q4-2 template. Put one of these in results/: "
        "result4-2_template.xlsx, result4-2.xlsx, result4-2(1).xlsx"
    )


def _capture_row_template(ws, row: int, max_col: int) -> dict:
    cells = []
    for col in range(1, max_col + 1):
        c = ws.cell(row, col)
        cells.append({
            "style": copy(c._style),
            "number_format": c.number_format,
        })
    return {
        "cells": cells,
        "height": ws.row_dimensions[row].height,
    }


def _apply_row_template(ws, row: int, template: dict) -> None:
    for col, style_info in enumerate(template["cells"], start=1):
        c = ws.cell(row, col)
        c._style = copy(style_info["style"])
        c.number_format = style_info["number_format"]
    if template["height"] is not None:
        ws.row_dimensions[row].height = template["height"]


def _validate_plan_sheet_dates(ws, dates: list[pd.Timestamp]) -> None:
    if ws.max_row != 335 or ws.max_column != 147:
        raise ValueError(
            "计划购电量 template shape changed; expected 335 rows x 147 columns."
        )

    if ws.cell(1, 2).value != "0:10-0:20":
        raise ValueError("Unexpected first interval header in 计划购电量 template.")
    if ws.cell(1, 145).value != "0:00-0:10+1":
        raise ValueError("Unexpected last interval header in 计划购电量 template.")

    for r, date in enumerate(dates, start=2):
        template_date = pd.Timestamp(ws.cell(r, 1).value).normalize()
        if template_date != date.normalize():
            raise ValueError(
                f"计划购电量 row {r} date mismatch: "
                f"{template_date.date()} != {date.date()}"
            )


def write_result4_2(
    plan: pd.DataFrame,
    dispatch: pd.DataFrame,
    emergency: pd.DataFrame,
    dates: list[pd.Timestamp],
) -> Path:
    """Fill official Q4-2 workbook from the FINAL rolling stochastic strategy."""
    template = find_template()
    wb = load_workbook(template)

    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    if wb.sheetnames != expected_sheets:
        raise ValueError(
            "The workbook is not the official Q4-2 template. "
            f"Found sheets: {wb.sheetnames}"
        )

    if len(dates) != 334:
        raise ValueError("Official result4-2.xlsx requires all 334 Feb-Dec days.")

    normalized_dates = [pd.Timestamp(d).normalize() for d in dates]

    # --------------------------------------------------------------
    # Sheet 1 — 计划购电量
    # Same official display rotation as Q2 template:
    # model 0..143 -> display 1..143,0.
    # --------------------------------------------------------------
    ws = wb["计划购电量"]
    _validate_plan_sheet_dates(ws, normalized_dates)

    for r, date in enumerate(normalized_dates, start=2):
        day = plan.loc[plan["date"] == date].sort_values("time_index")
        if len(day) != N:
            raise ValueError(f"{date.date()}: final plan does not have 144 slots.")

        values = day["grid_purchase_kwh"].to_numpy(float)
        display_values = np.concatenate([values[1:], values[:1]])
        for t, value in enumerate(display_values):
            ws.cell(r, 2 + t).value = float(value)

        ws.cell(r, 146).value = float(values.sum())
        ws.cell(r, 147).value = float(day["grid_cost_yuan"].sum())

    # --------------------------------------------------------------
    # Sheet 2 — 充放电量
    # --------------------------------------------------------------
    ws = wb["充放电量"]
    if ws.max_column != 6 or ws.max_row < 7:
        raise ValueError("Unexpected 充放电量 template structure.")

    block_templates = [
        _capture_row_template(ws, row, 6) for row in range(2, 8)
    ]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    periods = [
        "0:00-4:00", "4:00-8:00", "8:00-12:00",
        "12:00-16:00", "16:00-20:00", "20:00-24:00",
    ]

    out_row = 2
    for date in normalized_dates:
        day = dispatch.loc[dispatch["date"] == date].sort_values("time_index")
        if len(day) != N:
            raise ValueError(f"{date.date()}: dispatch does not have 144 slots.")

        for block in range(6):
            r = out_row + block
            _apply_row_template(ws, r, block_templates[block])
            part = day.iloc[24 * block: 24 * (block + 1)]

            ws.cell(r, 1).value = date.to_pydatetime() if block == 0 else None
            ws.cell(r, 2).value = periods[block]
            ws.cell(r, 3).value = float(part["charge_kwh"].sum())
            ws.cell(r, 4).value = float(part["discharge_kwh"].sum())

            if block == 0:
                ws.cell(r, 5).value = "0:00"
                ws.cell(r, 6).value = float(day["soc_start_kwh"].iloc[0])
            elif block == 1:
                ws.cell(r, 5).value = "24:00"
                ws.cell(r, 6).value = float(day["soc_end_kwh"].iloc[-1])
            else:
                ws.cell(r, 5).value = None
                ws.cell(r, 6).value = None

        out_row += 6

    # --------------------------------------------------------------
    # Sheet 3 — 紧急购电量
    # --------------------------------------------------------------
    ws = wb["紧急购电量"]
    if ws.max_column != 3 or ws.max_row < 2:
        raise ValueError("Unexpected 紧急购电量 template structure.")

    row_template = _capture_row_template(ws, 2, 3)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    compressed = compress_emergency(emergency)
    out_row = 2

    if not compressed.empty:
        for date, group in compressed.groupby("date", sort=True):
            group = group.sort_values("start_index").reset_index(drop=True)
            for j, rec in group.iterrows():
                _apply_row_template(ws, out_row, row_template)
                ws.cell(out_row, 1).value = (
                    pd.Timestamp(date).to_pydatetime() if j == 0 else None
                )
                ws.cell(out_row, 2).value = str(rec["period"])
                ws.cell(out_row, 3).value = float(rec["emergency_kwh"])
                out_row += 1

    output = ROOT / "results" / "result4-2.xlsx"
    wb.save(output)
    return output


# ======================================================================
# Main
# ======================================================================
def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    price_df, actual, forecast = read_inputs()

    dates = list(DATES)
    if MAX_DAYS > 0:
        dates = dates[:MAX_DAYS]

    all_results = {}

    # Run both steps separately because their causal SOC chains can diverge.
    for strategy in ["dynamic_greedy", "dynamic_rolling"]:
        feb1_soc, january_residuals, january_warmup = warm_up_january(
            price_df, actual, forecast, strategy
        )

        plan, dispatch, emergency, rolling_detail, daily = run_strategy(
            name=strategy,
            price_df=price_df,
            actual=actual,
            forecast=forecast,
            initial_soc=feb1_soc,
            january_residuals=january_residuals,
            dates=dates,
            keep_detail=True,
        )

        all_results[strategy] = {
            "feb1_soc": feb1_soc,
            "january_warmup": january_warmup,
            "plan": plan,
            "dispatch": dispatch,
            "emergency": emergency,
            "rolling_detail": rolling_detail,
            "daily": daily,
            "stats": summarize_strategy(daily),
        }

        january_warmup.to_csv(
            OUT / f"{strategy}_january_warmup.csv",
            index=False,
            encoding="utf-8-sig",
        )
        daily.to_csv(
            OUT / f"{strategy}_daily_summary.csv",
            index=False,
            encoding="utf-8-sig",
        )
        plan.to_csv(
            OUT / f"{strategy}_plan.csv",
            index=False,
            encoding="utf-8-sig",
        )
        dispatch.to_csv(
            OUT / f"{strategy}_dispatch.csv",
            index=False,
            encoding="utf-8-sig",
        )
        emergency.to_csv(
            OUT / f"{strategy}_emergency.csv",
            index=False,
            encoding="utf-8-sig",
        )
        rolling_detail.to_csv(
            OUT / f"{strategy}_rolling_diagnostics.csv",
            index=False,
            encoding="utf-8-sig",
        )

    comparison = pd.DataFrame([
        all_results["dynamic_greedy"]["stats"],
        all_results["dynamic_rolling"]["stats"],
    ])

    baseline_cost = float(
        comparison.loc[
            comparison["strategy"] == "dynamic_greedy",
            "total_cost_yuan",
        ].iloc[0]
    )
    comparison["saving_vs_dynamic_greedy_yuan"] = (
        baseline_cost - comparison["total_cost_yuan"]
    )
    comparison["saving_vs_dynamic_greedy_pct"] = (
        100.0 * comparison["saving_vs_dynamic_greedy_yuan"] / baseline_cost
    )
    comparison.to_csv(
        OUT / "strategy_comparison.csv",
        index=False,
        encoding="utf-8-sig",
    )

    # Official workbook uses Step 2 / final rolling strategy.
    final = all_results[FINAL_STRATEGY_NAME]
    if len(dates) == len(DATES):
        try:
            excel = write_result4_2(
                final["plan"], final["dispatch"], final["emergency"], dates
            )
        except FileNotFoundError:
            excel = None
    else:
        excel = None

    # Continuity / feasibility diagnostics for final strategy.
    final_daily = final["daily"]
    continuity = (
        final_daily["soc_end_kwh"].iloc[:-1].to_numpy()
        - final_daily["soc_start_kwh"].iloc[1:].to_numpy()
        if len(final_daily) > 1 else np.array([0.0])
    )
    terminal = final_daily["soc_end_kwh"].to_numpy(float)

    b = all_results["dynamic_greedy"]["stats"]
    f = all_results["dynamic_rolling"]["stats"]
    saving_yuan = b["total_cost_yuan"] - f["total_cost_yuan"]
    saving_pct = 100.0 * saving_yuan / b["total_cost_yuan"]

    summary = f"""Q4-2 DYNAMIC PRICE SUMMARY
============================
Price input:                    Attachment 4, date-specific 144-slot prices
Forecast model:                 unchanged from Q2
Scenario generation:            unchanged from Q2
Safety scenario quantile:       {SAFETY_QUANTILE:.2f}
CVaR:                           NOT used
Day-ahead decision:             planned grid purchase fixed at 0:00

STEP 1 — dynamic_greedy
-----------------------
Total cost:                     {b['total_cost_yuan']:,.2f} yuan
Planned purchase cost:          {b['planned_cost_yuan']:,.2f} yuan
Emergency cost:                 {b['emergency_cost_yuan']:,.2f} yuan
Emergency energy:               {b['emergency_purchase_kwh']:,.2f} kWh
Emergency days:                 {b['emergency_days']} / {len(final_daily)}
Emergency intervals:            {b['emergency_intervals']}

STEP 2 — dynamic_rolling (FINAL)
---------------------------------------
Total cost:                     {f['total_cost_yuan']:,.2f} yuan
Planned purchase cost:          {f['planned_cost_yuan']:,.2f} yuan
Emergency cost:                 {f['emergency_cost_yuan']:,.2f} yuan
Emergency energy:               {f['emergency_purchase_kwh']:,.2f} kWh
Emergency days:                 {f['emergency_days']} / {len(final_daily)}
Emergency intervals:            {f['emergency_intervals']}
Rolling LP calls:               {f['rolling_lp_calls']}

Step-2 vs Step-1 economic change
--------------------------------
Cost saving:                    {saving_yuan:,.2f} yuan
Cost saving rate:               {saving_pct:.3f}%

Final battery / feasibility
---------------------------
Jan-1 initial SOC:              {SOC_JAN1:,.2f} kWh
Feb-1 initial SOC:              {final['feb1_soc']:,.2f} kWh
Terminal SOC range:             {terminal.min():,.2f} to {terminal.max():,.2f} kWh
SOC continuity error:           {np.max(np.abs(continuity)):.3e} kWh
Balance error:                  {final_daily.attrs['max_balance_error']:.3e} kWh

Official workbook
-----------------
{excel if excel is not None else 'Not written (template absent or debug mode)'}
"""

    (OUT / "question4_2_summary.txt").write_text(summary, encoding="utf-8")
    print(summary)
    print("\nStrategy comparison:\n", comparison.to_string(index=False))


if __name__ == "__main__":
    main()
