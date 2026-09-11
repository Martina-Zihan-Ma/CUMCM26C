#!/usr/bin/env python3
"""
Q2 Step 2 — day-ahead grid planning + causal real-time battery dispatch.

This script consumes dynamic_forecasts.csv produced by q2_forecast.py.
The forecasting framework is unchanged.

Planning layer
--------------
1. Historical forecast residuals are bootstrap-sampled into discrete scenarios.
2. The day-ahead decision is the planned grid-purchase vector G_t.
3. Battery charge/discharge is scenario recourse used to value storage flexibility.
4. The scenario at the SAFETY_QUANTILE by daily positive net-demand energy is
   required to have zero emergency purchase.
5. The objective is planned grid cost + expected 5x emergency-purchase cost.
6. No terminal reserve is imposed.

Execution layer
---------------
The planned grid purchase is fixed.  Actual battery operation is strictly causal:
- surplus grid+PV is charged into the battery up to power/capacity limits, then spilled;
- a deficit is supplied by battery first up to power/SOC limits;
- only the remaining deficit is emergency-purchased at 5x price.

Thus small forecast errors are absorbed by storage whenever physically possible,
rather than automatically becoming emergency purchases.

SOC initialization
------------------
SOC(2025-01-01 00:00) = 6000 kWh.  Jan 1 is kept as the initialization day.
From Jan 2 onward, actual causal terminal SOC becomes the next day's initial SOC.
"""

from __future__ import annotations

import os
from copy import copy, deepcopy
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


ROOT = Path(__file__).resolve().parents[1]
PRICE_FILE = ROOT / "data" / "processed" / "attachment1_standard_day.csv"
ACTUAL_FILE = ROOT / "data" / "processed" / "attachment2_actual_long.csv"
FORECAST_FILE = ROOT / "outputs" / "question2" / "dynamic_forecasts.csv"
OUT = ROOT / "outputs" / "question2"

JAN1 = pd.Timestamp("2025-01-01")
WARMUP_START = pd.Timestamp("2025-01-02")
WARMUP_END = pd.Timestamp("2025-01-31")
START = pd.Timestamp("2025-02-01")
END = pd.Timestamp("2025-12-31")
DATES = pd.date_range(START, END, freq="D")

N = 144
DT = 1.0 / 6.0

ETA_C = 0.90
ETA_D = 0.90
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_JAN1 = 6000.0
MAX_INTERVAL = 5000.0 * DT

EMERGENCY_MULTIPLIER = 5.0

# Scenario uncertainty.
RESIDUAL_DAYS = 30
RESIDUAL_HALF_LIFE = 14.0
SCENARIOS = 20
MIN_RESIDUAL_DAYS = 5
RNG_SEED = 20260911

# Scenario safety rule.
# The sampled scenario at this quantile of daily positive net-demand energy
# must be feasible with zero emergency purchase.
SAFETY_QUANTILE = 0.80


# Tiny penalties only break degenerate LP solutions.
CYCLE_EPS = 1e-6
SPILL_EPS = 1e-8

# Optional debug run, e.g. Q2_MAX_DAYS=5 python src/q2_optimize_final_clean.py
MAX_DAYS = int(os.environ.get("Q2_MAX_DAYS", "0"))


# ----------------------------------------------------------------------
# Input
# ----------------------------------------------------------------------

def read_inputs() -> tuple[np.ndarray, pd.DataFrame, pd.DataFrame]:
    for path in [PRICE_FILE, ACTUAL_FILE, FORECAST_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    price_df = pd.read_csv(PRICE_FILE).sort_values("time_index")
    actual = pd.read_csv(ACTUAL_FILE)
    forecast = pd.read_csv(FORECAST_FILE)

    actual["date"] = pd.to_datetime(actual["date"]).dt.normalize()
    forecast["date"] = pd.to_datetime(forecast["date"]).dt.normalize()

    if "history_end_date" in forecast.columns:
        forecast["history_end_date"] = pd.to_datetime(
            forecast["history_end_date"]
        ).dt.normalize()

    if price_df["time_index"].astype(int).tolist() != list(range(N)):
        raise ValueError("Price time_index must be 0..143.")

    required_actual = {"date", "time_index", "load_kw", "pv_actual_kw"}
    required_forecast = {
        "date", "time_index", "forecast_load_kw", "forecast_pv_kw",
        "load_residual_kw", "pv_residual_kw",
    }
    missing_actual = required_actual - set(actual.columns)
    missing_forecast = required_forecast - set(forecast.columns)
    if missing_actual:
        raise ValueError(f"Actual data missing: {sorted(missing_actual)}")
    if missing_forecast:
        raise ValueError(
            "Run q2_forecast_final_clean.py first. Missing forecast columns: "
            f"{sorted(missing_forecast)}"
        )

    for name, df in [("actual", actual), ("forecast", forecast)]:
        for date, part in df.groupby("date"):
            idx = part.sort_values("time_index")["time_index"].astype(int).tolist()
            if idx != list(range(N)):
                raise ValueError(f"{name} {date.date()}: time_index must be 0..143.")

    if forecast["date"].min() > WARMUP_START:
        raise ValueError("Forecast file must include January warm-up forecasts from Jan 2.")
    if forecast["date"].max() < END:
        raise ValueError("Forecast file must include forecasts through Dec 31.")

    if "history_end_date" in forecast.columns:
        if not (forecast["history_end_date"] < forecast["date"]).all():
            raise RuntimeError("Forecast leakage: history_end_date >= target date.")

    return price_df["price_yuan_per_kwh"].to_numpy(float), actual, forecast


def day_values(df: pd.DataFrame, date: pd.Timestamp, col: str) -> np.ndarray:
    x = df.loc[df["date"] == date].sort_values("time_index")
    if len(x) != N:
        raise ValueError(f"{date.date()}: {col} has {len(x)} rows, expected {N}.")
    return x[col].to_numpy(float)


def proposed_forecast(
    forecast: pd.DataFrame,
    date: pd.Timestamp,
) -> tuple[np.ndarray, np.ndarray]:
    return (
        day_values(forecast, date, "forecast_load_kw"),
        day_values(forecast, date, "forecast_pv_kw"),
    )


# ----------------------------------------------------------------------
# Unified forecast residual history
# ----------------------------------------------------------------------

def positive_net_energy(load_kw: np.ndarray, pv_kw: np.ndarray) -> float:
    return float(np.maximum(load_kw - pv_kw, 0.0).sum() * DT)


def forecast_error_risk(
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
) -> float:
    """Positive actual-minus-forecast net-demand error, integrated over a day."""
    net_error_kw = (load_actual - load_hat) - (pv_actual - pv_hat)
    return float(np.maximum(net_error_kw, 0.0).sum() * DT)


def residual_record(
    date: pd.Timestamp,
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
) -> dict:
    return {
        "date": date.normalize(),
        "load_error_kw": (load_actual - load_hat).astype(float),
        "pv_error_kw": (pv_actual - pv_hat).astype(float),
        "forecast_net_kwh": positive_net_energy(load_hat, pv_hat),
        "risk_kwh": forecast_error_risk(
            load_hat, pv_hat, load_actual, pv_actual
        ),
    }


def recent_residuals(history: list[dict], date: pd.Timestamp) -> list[dict]:
    prior = sorted(
        [r for r in history if r["date"] < date],
        key=lambda r: r["date"],
    )
    return prior[-RESIDUAL_DAYS:]


# ----------------------------------------------------------------------
# Residual bootstrap scenarios
# ----------------------------------------------------------------------

def sample_scenarios(
    date: pd.Timestamp,
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    residual_history: list[dict],
    rng: np.random.Generator,
) -> tuple[np.ndarray, np.ndarray]:
    residuals = recent_residuals(residual_history, date)

    if len(residuals) < MIN_RESIDUAL_DAYS:
        return (
            np.maximum(load_hat, 0.0)[None, :] * DT,
            np.maximum(pv_hat, 0.0)[None, :] * DT,
        )

    ages = np.array([(date - r["date"]).days for r in residuals], dtype=float)
    weights = np.exp(-np.log(2.0) * ages / RESIDUAL_HALF_LIFE)
    weights /= weights.sum()

    selected = rng.choice(
        len(residuals), size=SCENARIOS, replace=True, p=weights
    )

    load_scenarios = []
    pv_scenarios = []
    for i in selected:
        r = residuals[int(i)]
        load_scenarios.append(
            np.maximum(load_hat + r["load_error_kw"], 0.0) * DT
        )
        pv_scenarios.append(
            np.maximum(pv_hat + r["pv_error_kw"], 0.0) * DT
        )

    return np.vstack(load_scenarios), np.vstack(pv_scenarios)


# ----------------------------------------------------------------------
# Scenario-based day-ahead stochastic linear program
# ----------------------------------------------------------------------

@dataclass
class VarIndex:
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    soc: np.ndarray
    emergency: np.ndarray
    spill: np.ndarray


def make_index(k: int) -> tuple[VarIndex, int]:
    p = 0
    grid = np.arange(p, p + N); p += N
    charge = np.arange(p, p + k * N).reshape(k, N); p += k * N
    discharge = np.arange(p, p + k * N).reshape(k, N); p += k * N
    soc = np.arange(p, p + k * (N + 1)).reshape(k, N + 1); p += k * (N + 1)
    emergency = np.arange(p, p + k * N).reshape(k, N); p += k * N
    spill = np.arange(p, p + k * N).reshape(k, N); p += k * N
    return VarIndex(grid, charge, discharge, soc, emergency, spill), p


def choose_safety_scenario(net_scenarios: np.ndarray) -> tuple[int, float]:
    """
    Pick one sampled whole-day scenario at the requested quantile.

    Severity is the day's positive net-demand energy.  Using a whole historical
    residual path preserves the temporal correlation inside the sampled scenario.
    """
    severity = np.maximum(net_scenarios, 0.0).sum(axis=1)
    order = np.argsort(severity)
    rank = int(np.ceil(SAFETY_QUANTILE * len(order))) - 1
    rank = int(np.clip(rank, 0, len(order) - 1))
    idx = int(order[rank])
    return idx, float(severity[idx])


def solve_plan(
    price: np.ndarray,
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    initial_soc: float,
) -> dict:
    """
    Day-ahead plan.

    First-stage decision
        grid[t] : planned grid purchase, fixed at 0:00.

    Scenario recourse
        charge/discharge/SOC/emergency/spill may adapt to each sampled scenario.

    Safety rule
        The sampled scenario at SAFETY_QUANTILE is not allowed emergency purchase.

    Objective
        planned grid cost + expected 5x emergency cost.
    """
    k = load_scenarios.shape[0]
    if load_scenarios.shape != (k, N) or pv_scenarios.shape != (k, N):
        raise ValueError("Bad scenario shape.")

    net_scenarios = load_scenarios - pv_scenarios
    safety_idx, safety_energy = choose_safety_scenario(net_scenarios)

    ix, nv = make_index(k)

    c = np.zeros(nv)
    c[ix.grid] = price
    for w in range(k):
        c[ix.charge[w]] = CYCLE_EPS / k
        c[ix.discharge[w]] = CYCLE_EPS / k
        c[ix.emergency[w]] = EMERGENCY_MULTIPLIER * price / k
        c[ix.spill[w]] = SPILL_EPS / k

    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)

    for w in range(k):
        ub[ix.charge[w]] = MAX_INTERVAL
        ub[ix.discharge[w]] = MAX_INTERVAL
        lb[ix.soc[w]] = SOC_MIN
        ub[ix.soc[w]] = SOC_MAX
        lb[ix.soc[w, 0]] = initial_soc
        ub[ix.soc[w, 0]] = initial_soc

    # The selected safety scenario must be covered without emergency purchase.
    ub[ix.emergency[safety_idx]] = 0.0

    # Scenario energy balances + scenario SOC dynamics.
    Aeq = lil_matrix((k * N + k * N, nv), dtype=float)
    beq = np.zeros(2 * k * N, dtype=float)
    row = 0

    for w in range(k):
        for t in range(N):
            # G + E + D = (Load - PV) + C + Spill
            Aeq[row, ix.grid[t]] = 1.0
            Aeq[row, ix.emergency[w, t]] = 1.0
            Aeq[row, ix.discharge[w, t]] = 1.0
            Aeq[row, ix.charge[w, t]] = -1.0
            Aeq[row, ix.spill[w, t]] = -1.0
            beq[row] = net_scenarios[w, t]
            row += 1

    for w in range(k):
        for t in range(N):
            # SOC_{t+1} = SOC_t + eta_c*C_t - D_t/eta_d
            Aeq[row, ix.soc[w, t + 1]] = 1.0
            Aeq[row, ix.soc[w, t]] = -1.0
            Aeq[row, ix.charge[w, t]] = -ETA_C
            Aeq[row, ix.discharge[w, t]] = 1.0 / ETA_D
            row += 1

    # Combined charge/discharge power envelope in each scenario and interval.
    Aub = lil_matrix((k * N, nv), dtype=float)
    bub = np.full(k * N, MAX_INTERVAL, dtype=float)
    row = 0
    for w in range(k):
        for t in range(N):
            Aub[row, ix.charge[w, t]] = 1.0
            Aub[row, ix.discharge[w, t]] = 1.0
            row += 1

    result = linprog(
        c=c,
        A_ub=Aub.tocsr(),
        b_ub=bub,
        A_eq=Aeq.tocsr(),
        b_eq=beq,
        bounds=list(zip(lb, ub)),
        method="highs",
    )
    if not result.success:
        raise RuntimeError(result.message)

    x = result.x
    scenario_emergency_cost = np.array([
        np.dot(x[ix.emergency[w]], EMERGENCY_MULTIPLIER * price)
        for w in range(k)
    ])

    return {
        "grid": x[ix.grid].copy(),
        "expected_emergency_cost": float(scenario_emergency_cost.mean()),
        "safety_scenario_index": safety_idx,
        "safety_scenario_positive_net_kwh": safety_energy,
        "safety_scenario_emergency_kwh": float(x[ix.emergency[safety_idx]].sum()),
    }


def execute_causal_dispatch(
    grid: np.ndarray,
    initial_soc: float,
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
) -> dict:
    """
    Strictly causal real-time battery policy.

    At interval t only current grid purchase, actual Load/PV and current SOC are used.
    No future actual data are used.
    """
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    spill = np.zeros(N)
    soc_start = np.zeros(N)
    soc_end = np.zeros(N)

    soc = float(initial_soc)

    for t in range(N):
        soc_start[t] = soc
        balance = float(grid[t] + pv_kwh[t] - load_kwh[t])

        if balance >= 0.0:
            # Surplus: charge first; only unavoidable remainder is spilled.
            capacity_input = max((SOC_MAX - soc) / ETA_C, 0.0)
            charge[t] = min(balance, MAX_INTERVAL, capacity_input)
            soc += ETA_C * charge[t]
            spill[t] = max(balance - charge[t], 0.0)
        else:
            # Deficit: discharge first; emergency purchase is the last resort.
            deficit = -balance
            available_output = max((soc - SOC_MIN) * ETA_D, 0.0)
            discharge[t] = min(deficit, MAX_INTERVAL, available_output)
            soc -= discharge[t] / ETA_D
            emergency[t] = max(deficit - discharge[t], 0.0)

        soc = float(np.clip(soc, SOC_MIN, SOC_MAX))
        soc_end[t] = soc

    balance_error = np.max(np.abs(
        grid + emergency + pv_kwh + discharge
        - load_kwh - charge - spill
    ))

    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "soc_start": soc_start,
        "soc_end": soc_end,
        "terminal_soc": float(soc),
        "balance_error": float(balance_error),
    }


# ----------------------------------------------------------------------
# January SOC warm-up using the same forecasts and causal execution
# ----------------------------------------------------------------------

def warm_up_january(
    price: np.ndarray,
    actual: pd.DataFrame,
    forecast: pd.DataFrame,
) -> tuple[float, list[dict], pd.DataFrame]:
    rng = np.random.default_rng(RNG_SEED)

    # Explicit initialization assumption for Jan 1 only.
    soc = SOC_JAN1
    residuals: list[dict] = []
    rows = [{
        "date": JAN1,
        "soc_start_kwh": SOC_JAN1,
        "soc_end_kwh": SOC_JAN1,
        "note": "initialization day; no prior history",
    }]

    for date in pd.date_range(WARMUP_START, WARMUP_END, freq="D"):
        load_hat, pv_hat = proposed_forecast(forecast, date)
        load_s, pv_s = sample_scenarios(date, load_hat, pv_hat, residuals, rng)
        plan = solve_plan(price, load_s, pv_s, soc)

        load_actual_kw = day_values(actual, date, "load_kw")
        pv_actual_kw = day_values(actual, date, "pv_actual_kw")
        dispatch = execute_causal_dispatch(
            plan["grid"],
            soc,
            load_actual_kw * DT,
            pv_actual_kw * DT,
        )

        rows.append({
            "date": date,
            "soc_start_kwh": soc,
            "soc_end_kwh": dispatch["terminal_soc"],
            "planned_purchase_kwh": float(plan["grid"].sum()),
            "actual_emergency_kwh": float(dispatch["emergency"].sum()),
            "note": "causal battery warm-up",
        })

        # Only after the day finishes may its actual error affect future scenarios.
        residuals.append(residual_record(
            date, load_hat, pv_hat, load_actual_kw, pv_actual_kw
        ))
        soc = dispatch["terminal_soc"]

    return float(soc), residuals, pd.DataFrame(rows)


# ----------------------------------------------------------------------
# Feb-Dec strategy
# ----------------------------------------------------------------------

def run_optimized_strategy(
    name: str,
    price: np.ndarray,
    actual: pd.DataFrame,
    forecast: pd.DataFrame,
    initial_soc: float,
    january_residuals: list[dict],
    dates: list[pd.Timestamp],
    keep_detail: bool,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rng = np.random.default_rng(RNG_SEED + 1)
    residuals = deepcopy(january_residuals)
    soc = float(initial_soc)

    plan_rows: list[dict] = []
    dispatch_rows: list[dict] = []
    emergency_rows: list[dict] = []
    daily_rows: list[dict] = []
    max_balance_error = 0.0

    for date in dates:
        load_hat, pv_hat = proposed_forecast(forecast, date)
        load_s, pv_s = sample_scenarios(date, load_hat, pv_hat, residuals, rng)

        # Day-ahead: only planned grid purchase is binding in real operation.
        plan = solve_plan(price, load_s, pv_s, soc)

        # Real-time: actual Load/PV are handled by the causal battery policy.
        load_actual_kw = day_values(actual, date, "load_kw")
        pv_actual_kw = day_values(actual, date, "pv_actual_kw")
        dispatch = execute_causal_dispatch(
            plan["grid"],
            soc,
            load_actual_kw * DT,
            pv_actual_kw * DT,
        )

        emergency = dispatch["emergency"]
        spill = dispatch["spill"]
        max_balance_error = max(max_balance_error, dispatch["balance_error"])

        planned_cost = float(np.dot(plan["grid"], price))
        emergency_cost = float(np.dot(
            emergency, EMERGENCY_MULTIPLIER * price
        ))

        daily_rows.append({
            "strategy": name,
            "date": date,
            "soc_start_kwh": soc,
            "soc_end_kwh": dispatch["terminal_soc"],
            "planned_purchase_kwh": float(plan["grid"].sum()),
            "planned_cost_yuan": planned_cost,
            "expected_emergency_cost_yuan": plan["expected_emergency_cost"],
            "safety_quantile": SAFETY_QUANTILE,
            "safety_scenario_positive_net_kwh": (
                plan["safety_scenario_positive_net_kwh"]
            ),
            "actual_emergency_kwh": float(emergency.sum()),
            "actual_emergency_cost_yuan": emergency_cost,
            "actual_emergency_intervals": int((emergency > 1e-8).sum()),
            "spill_kwh": float(spill.sum()),
            "total_cost_yuan": planned_cost + emergency_cost,
        })

        if keep_detail:
            for t in range(N):
                plan_rows.append({
                    "date": date,
                    "time_index": t,
                    "grid_purchase_kwh": float(plan["grid"][t]),
                    "price_yuan_per_kwh": float(price[t]),
                    "grid_cost_yuan": float(plan["grid"][t] * price[t]),
                })
                dispatch_rows.append({
                    "date": date,
                    "time_index": t,
                    "charge_kwh": float(dispatch["charge"][t]),
                    "discharge_kwh": float(dispatch["discharge"][t]),
                    "soc_start_kwh": float(dispatch["soc_start"][t]),
                    "soc_end_kwh": float(dispatch["soc_end"][t]),
                    "spill_kwh": float(spill[t]),
                })
                emergency_rows.append({
                    "date": date,
                    "time_index": t,
                    "emergency_purchase_kwh": float(emergency[t]),
                    "emergency_cost_yuan": float(
                        emergency[t] * EMERGENCY_MULTIPLIER * price[t]
                    ),
                })

        # Update future uncertainty only after today's realized values are observed.
        residuals.append(residual_record(
            date, load_hat, pv_hat, load_actual_kw, pv_actual_kw
        ))
        soc = dispatch["terminal_soc"]

    daily = pd.DataFrame(daily_rows)
    daily.attrs["max_balance_error"] = max_balance_error

    return (
        pd.DataFrame(plan_rows),
        pd.DataFrame(dispatch_rows),
        pd.DataFrame(emergency_rows),
        daily,
    )


def no_storage_same_forecast(
    price: np.ndarray,
    actual: pd.DataFrame,
    forecast: pd.DataFrame,
    dates: list[pd.Timestamp],
) -> pd.DataFrame:
    rows = []
    for date in dates:
        load_hat, pv_hat = proposed_forecast(forecast, date)
        load_actual = day_values(actual, date, "load_kw") * DT
        pv_actual = day_values(actual, date, "pv_actual_kw") * DT

        grid = np.maximum(load_hat - pv_hat, 0.0) * DT
        balance = grid + pv_actual - load_actual
        emergency = np.maximum(-balance, 0.0)
        spill = np.maximum(balance, 0.0)

        planned_cost = float(np.dot(grid, price))
        emergency_cost = float(np.dot(
            emergency, EMERGENCY_MULTIPLIER * price
        ))

        rows.append({
            "strategy": "no_storage_same_forecast",
            "date": date,
            "planned_purchase_kwh": float(grid.sum()),
            "planned_cost_yuan": planned_cost,
            "actual_emergency_kwh": float(emergency.sum()),
            "actual_emergency_cost_yuan": emergency_cost,
            "actual_emergency_intervals": int((emergency > 1e-8).sum()),
            "spill_kwh": float(spill.sum()),
            "total_cost_yuan": planned_cost + emergency_cost,
        })
    return pd.DataFrame(rows)


def summarize_strategy(daily: pd.DataFrame) -> dict:
    emergency_kwh = float(daily["actual_emergency_kwh"].sum())
    emergency_cost = float(daily["actual_emergency_cost_yuan"].sum())
    return {
        "strategy": daily["strategy"].iloc[0],
        "planned_purchase_kwh": float(daily["planned_purchase_kwh"].sum()),
        "planned_cost_yuan": float(daily["planned_cost_yuan"].sum()),
        "emergency_purchase_kwh": emergency_kwh,
        "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": float(daily["total_cost_yuan"].sum()),
        "spill_kwh": float(daily["spill_kwh"].sum()),
        "emergency_days": int((daily["actual_emergency_kwh"] > 1e-8).sum()),
        "emergency_intervals": int(daily["actual_emergency_intervals"].sum()),
        "avg_normal_price_of_emergency": (
            emergency_cost / (EMERGENCY_MULTIPLIER * emergency_kwh)
            if emergency_kwh > 1e-12 else np.nan
        ),
    }

# ----------------------------------------------------------------------
# Official workbook for the final strategy
# ----------------------------------------------------------------------

def _format_template_time(total_minutes: int) -> str:
    """
    Format a time boundary using the convention in the official result2 template.

    time_index=0 represents 0:00-0:10 and time_index=143 represents
    23:50-24:00.  Boundaries at/after 24:00 receive the '+1' suffix.
    """
    day_offset, minute = divmod(int(total_minutes), 24 * 60)
    hour, minute = divmod(minute, 60)
    label = f"{hour}:{minute:02d}"
    if day_offset > 0:
        label += f"+{day_offset}"
    return label


def compress_emergency(emergency: pd.DataFrame) -> pd.DataFrame:
    """
    Merge consecutive 10-minute emergency-purchase slots into time periods.

    IMPORTANT:
    Each time_index denotes the preceding 10-minute interval:
        time_index 0   -> 0:00-0:10
        ...
        time_index 143 -> 23:50-0:00+1

    This is consistent with the source timestamps, where index 0 is labelled 00:10.
    """
    rows = []
    active = emergency.loc[emergency["emergency_purchase_kwh"] > 1e-8].copy()
    if active.empty:
        return pd.DataFrame(
            columns=["date", "start_index", "end_index", "period", "emergency_kwh"]
        )

    for date, day in active.groupby("date"):
        day = day.sort_values("time_index")
        idx = day["time_index"].astype(int).tolist()
        values = day["emergency_purchase_kwh"].to_numpy(float)

        start = 0
        for j in range(1, len(idx) + 1):
            end_group = j == len(idx) or idx[j] != idx[j - 1] + 1
            if not end_group:
                continue

            a = idx[start]
            last = idx[j - 1]
            # index t is the interval [t*10, (t+1)*10] minutes.
            start_min = a * 10
            end_min = (last + 1) * 10

            rows.append({
                "date": pd.Timestamp(date).normalize(),
                "start_index": int(a),
                "end_index": int(last),
                "period": (
                    f"{_format_template_time(start_min)}-"
                    f"{_format_template_time(end_min)}"
                ),
                "emergency_kwh": float(values[start:j].sum()),
            })
            start = j

    return pd.DataFrame(rows)


def find_template() -> Path:
    """Locate the official blank result2 workbook."""
    candidates = [
        ROOT / "results" / "result2_template.xlsx",
        ROOT / "results" / "result2_template(1).xlsx",
        ROOT / "results" / "result2.xlsx",
        ROOT / "result2_template.xlsx",
        ROOT / "result2_template(1).xlsx",
        ROOT / "result2.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        "Cannot find the official result2 template. "
        "Put result2_template.xlsx under the project results/ folder."
    )


def _capture_row_template(ws, row: int, max_col: int) -> dict:
    """Capture style/height from one template row before rows are rebuilt."""
    cells = []
    for col in range(1, max_col + 1):
        c = ws.cell(row, col)
        cells.append({
            "style": copy(c._style),
            "number_format": c.number_format,
        })
    return {
        "cells": cells,
        "height": ws.row_dimensions[row].height,
    }


def _apply_row_template(ws, row: int, template: dict) -> None:
    """Apply a previously captured template row style to a newly created row."""
    for col, style_info in enumerate(template["cells"], start=1):
        c = ws.cell(row, col)
        c._style = copy(style_info["style"])
        c.number_format = style_info["number_format"]
    if template["height"] is not None:
        ws.row_dimensions[row].height = template["height"]


def _validate_plan_sheet_dates(ws, dates: list[pd.Timestamp]) -> None:
    if ws.max_row != 335 or ws.max_column != 147:
        raise ValueError(
            "计划购电量 template shape changed; expected 335 rows x 147 columns."
        )

    for r, date in enumerate(dates, start=2):
        template_date = pd.Timestamp(ws.cell(r, 1).value).normalize()
        if template_date != date.normalize():
            raise ValueError(
                f"计划购电量 row {r} date mismatch: "
                f"{template_date.date()} != {date.date()}"
            )


def write_result2(
    plan: pd.DataFrame,
    dispatch: pd.DataFrame,
    emergency: pd.DataFrame,
    dates: list[pd.Timestamp],
) -> Path:
    """
    Fill the official result2 workbook.

    Unlike the previous version, the ellipsis rows in Sheets 2 and 3 are treated
    as compact examples, NOT as a request to omit the middle dates.

    Sheet 1 — 计划购电量
        The official template already contains all 334 dates. Fill the 144
        10-minute planned purchases, daily total energy, and daily cost.

    Sheet 2 — 充放电量
        Expand the compact Feb-1 / Feb-2 / ... / Dec-31 example into all
        334 dates. Each date receives six 4-hour blocks plus 0:00/24:00 SOC.

    Sheet 3 — 紧急购电量
        Expand to all 334 dates. Consecutive emergency slots are merged.
        A day with no emergency purchase is written explicitly as "无", 0.

    Output:
        outputs/question2/result2_final.xlsx
    """
    template = find_template()
    wb = load_workbook(template)

    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    if wb.sheetnames != expected_sheets:
        raise ValueError(
            "The workbook is not the official result2 template. "
            f"Found sheets: {wb.sheetnames}"
        )

    if len(dates) != 334:
        raise ValueError("Official result2.xlsx requires all 334 Feb-Dec days.")

    normalized_dates = [pd.Timestamp(d).normalize() for d in dates]

    # ==============================================================
    # Sheet 1 — 计划购电量
    # ==============================================================
    ws = wb["计划购电量"]
    _validate_plan_sheet_dates(ws, normalized_dates)

    for r, date in enumerate(normalized_dates, start=2):
        day = plan.loc[plan["date"] == date].sort_values("time_index")
        if len(day) != N:
            raise ValueError(f"{date.date()}: final plan does not have 144 slots.")

        values = day["grid_purchase_kwh"].to_numpy(float)

        # B:EO = 144 interval purchases; EP = daily total; EQ = daily cost.
        for t, value in enumerate(values):
            ws.cell(r, 2 + t).value = float(value)

        ws.cell(r, 146).value = float(values.sum())
        ws.cell(r, 147).value = float(day["grid_cost_yuan"].sum())

    # ==============================================================
    # Sheet 2 — 充放电量
    # Expand the compact example to every Feb-Dec date.
    # ==============================================================
    ws = wb["充放电量"]
    if ws.max_column != 6 or ws.max_row < 7:
        raise ValueError("Unexpected 充放电量 template structure.")

    # Capture the six official row styles from the first sample block.
    block_templates = [
        _capture_row_template(ws, row, 6) for row in range(2, 8)
    ]

    # Rebuild all data rows while preserving the original header.
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    periods = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]

    out_row = 2
    for date in normalized_dates:
        day = dispatch.loc[dispatch["date"] == date].sort_values("time_index")
        if len(day) != N:
            raise ValueError(f"{date.date()}: dispatch does not have 144 slots.")

        for block in range(6):
            r = out_row + block
            _apply_row_template(ws, r, block_templates[block])

            part = day.iloc[24 * block: 24 * (block + 1)]

            ws.cell(r, 1).value = (
                date.to_pydatetime() if block == 0 else None
            )
            ws.cell(r, 2).value = periods[block]
            ws.cell(r, 3).value = float(part["charge_kwh"].sum())
            ws.cell(r, 4).value = float(part["discharge_kwh"].sum())

            if block == 0:
                ws.cell(r, 5).value = "0:00"
                ws.cell(r, 6).value = float(day["soc_start_kwh"].iloc[0])
            elif block == 1:
                ws.cell(r, 5).value = "24:00"
                ws.cell(r, 6).value = float(day["soc_end_kwh"].iloc[-1])
            else:
                ws.cell(r, 5).value = None
                ws.cell(r, 6).value = None

        out_row += 6

    # ==============================================================
    # Sheet 3 — 紧急购电量
    # Expand the compact example to every Feb-Dec date.
    # ==============================================================
    ws = wb["紧急购电量"]
    if ws.max_column != 3 or ws.max_row < 2:
        raise ValueError("Unexpected 紧急购电量 template structure.")

    emergency_row_template = _capture_row_template(ws, 2, 3)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    compressed = compress_emergency(emergency)
    out_row = 2

    for date in normalized_dates:
        if compressed.empty:
            group = compressed
        else:
            group = (
                compressed.loc[compressed["date"] == date]
                .sort_values("start_index")
                .reset_index(drop=True)
            )

        if group.empty:
            _apply_row_template(ws, out_row, emergency_row_template)
            ws.cell(out_row, 1).value = date.to_pydatetime()
            ws.cell(out_row, 2).value = "无"
            ws.cell(out_row, 3).value = 0.0
            out_row += 1
            continue

        for j, rec in group.iterrows():
            _apply_row_template(ws, out_row, emergency_row_template)
            ws.cell(out_row, 1).value = (
                date.to_pydatetime() if j == 0 else None
            )
            ws.cell(out_row, 2).value = str(rec["period"])
            ws.cell(out_row, 3).value = float(rec["emergency_kwh"])
            out_row += 1

    output = OUT / "result2_final.xlsx"
    wb.save(output)
    return output


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    price, actual, forecast = read_inputs()

    dates = list(DATES)
    if MAX_DAYS > 0:
        dates = dates[:MAX_DAYS]

    # One continuous SOC chain: Jan 1 initialization -> Jan 2-31 warm-up -> Feb-Dec.
    feb1_soc, january_residuals, january_warmup = warm_up_january(
        price, actual, forecast
    )

    no_storage = no_storage_same_forecast(price, actual, forecast, dates)
    plan, dispatch, emergency, final_daily = run_optimized_strategy(
        name="optimized_proposed",
        price=price,
        actual=actual,
        forecast=forecast,
        initial_soc=feb1_soc,
        january_residuals=january_residuals,
        dates=dates,
        keep_detail=True,
    )

    baseline_stats = summarize_strategy(no_storage)
    final_stats = summarize_strategy(final_daily)

    comparison = pd.DataFrame([baseline_stats, final_stats])
    baseline_cost = baseline_stats["total_cost_yuan"]
    comparison["saving_vs_no_storage_yuan"] = (
        baseline_cost - comparison["total_cost_yuan"]
    )
    comparison["saving_vs_no_storage_pct"] = (
        100.0 * comparison["saving_vs_no_storage_yuan"] / baseline_cost
    )

    # January residual diagnostics come from the SAME forecast model.
    january_residual_diagnostics = pd.DataFrame([
        {
            "date": r["date"],
            "forecast_net_kwh": r["forecast_net_kwh"],
            "risk_kwh": r["risk_kwh"],
        }
        for r in january_residuals
    ])

    january_residual_diagnostics.to_csv(
        OUT / "january_residual_diagnostics.csv", index=False, encoding="utf-8-sig"
    )
    january_warmup.to_csv(
        OUT / "january_warmup.csv", index=False, encoding="utf-8-sig"
    )
    comparison.to_csv(
        OUT / "strategy_comparison.csv", index=False, encoding="utf-8-sig"
    )
    final_daily.to_csv(
        OUT / "question2_daily_summary.csv", index=False, encoding="utf-8-sig"
    )
    plan.to_csv(OUT / "question2_plan.csv", index=False, encoding="utf-8-sig")
    dispatch.to_csv(
        OUT / "question2_dispatch.csv", index=False, encoding="utf-8-sig"
    )
    emergency.to_csv(
        OUT / "question2_emergency.csv", index=False, encoding="utf-8-sig"
    )

    if len(dates) == len(DATES):
        try:
            excel = write_result2(plan, dispatch, emergency, dates)
        except FileNotFoundError:
            excel = None
    else:
        excel = None

    continuity = (
        final_daily["soc_end_kwh"].iloc[:-1].to_numpy()
        - final_daily["soc_start_kwh"].iloc[1:].to_numpy()
        if len(final_daily) > 1 else np.array([0.0])
    )
    terminal = final_daily["soc_end_kwh"].to_numpy(float)

    saving_yuan = baseline_stats["total_cost_yuan"] - final_stats["total_cost_yuan"]
    saving_pct = 100.0 * saving_yuan / baseline_stats["total_cost_yuan"]

    summary = f"""Q2 FINAL SUMMARY
================
Safety scenario quantile:       {SAFETY_QUANTILE:.2f}
Day-ahead decision:             planned grid purchase
Real-time battery control:      causal surplus-charge / deficit-discharge
Terminal reserve:               none

Final strategy
--------------
Total purchase cost:           {final_stats['total_cost_yuan']:,.2f} yuan
Planned purchase cost:         {final_stats['planned_cost_yuan']:,.2f} yuan
Emergency purchase cost:       {final_stats['emergency_cost_yuan']:,.2f} yuan
Emergency purchase energy:     {final_stats['emergency_purchase_kwh']:,.2f} kWh
Emergency days:                {final_stats['emergency_days']} / {len(final_daily)}
Emergency intervals:           {final_stats['emergency_intervals']}
Spill energy:                  {final_stats['spill_kwh']:,.2f} kWh

No-storage baseline
-------------------
Total purchase cost:           {baseline_stats['total_cost_yuan']:,.2f} yuan
Emergency purchase cost:       {baseline_stats['emergency_cost_yuan']:,.2f} yuan

Economic improvement
--------------------
Cost saving:                   {saving_yuan:,.2f} yuan
Cost saving rate:              {saving_pct:.2f}%

Battery / feasibility
---------------------
Jan-1 initial SOC:             {SOC_JAN1:,.2f} kWh
Feb-1 initial SOC:             {feb1_soc:,.2f} kWh
Terminal SOC range:            {terminal.min():,.2f} to {terminal.max():,.2f} kWh
SOC continuity error:          {np.max(np.abs(continuity)):.3e} kWh
Balance error:                 {final_daily.attrs['max_balance_error']:.3e} kWh

Output workbook
---------------
{excel if excel is not None else 'Not written (template absent or debug mode)'}
"""
    (OUT / "question2_summary.txt").write_text(summary, encoding="utf-8")
    print(summary)


if __name__ == "__main__":
    main()
